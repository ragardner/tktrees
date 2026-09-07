# SPDX-License-Identifier: AGPL-3.0-only
# Copyright (c) R. A. Gardner

from __future__ import annotations

import csv
import os
import tempfile
import unittest

from openpyxl import Workbook
from src.session import Session
from tests.fixtures import ANIMALS_TABLE

FMT5 = [
    ["Animals", "All animals"],
    ["", "Cats", "Cat family"],
    ["", "", "Lion", "Lion"],
]
FMT6 = [
    ["Animals", "All animals", "note-a"],
    ["", "Cats", "Cat family", "note-c"],
    ["", "", "Lion", "Lion", "note-l"],
]
FMT7 = [
    ["L1", "L2", "L3", "Name"],
    ["Animals", "", "", "All animals"],
    ["", "Cats", "", "Cat family"],
    ["", "", "Lion", "Lion"],
]


class TestLoadFormats(unittest.TestCase):
    def test_format_1_from_export_flat(self):
        s = Session()
        s.load_table(ANIMALS_TABLE, id_col=0, parent_cols=[1], fmt=0)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "flat.csv")
            self.assertTrue(s.export_flat(path, hier="Parent", overwrite=True)["ok"])
            with open(path, newline="", encoding="utf-8") as fh:
                rows = list(csv.reader(fh))
        hier_cols = [i for i, name in enumerate(rows[0]) if name.startswith("Parent_")]
        s2 = Session()
        out = s2.load_table(rows, parent_cols=hier_cols, fmt=1)
        self.assertTrue(out["ok"], out)
        self.assertEqual(out["result"]["format"], 1)
        self.assertEqual(sorted(s2.nodes), ["animals", "cats", "lion"])
        self.assertEqual(s2.nodes["lion"].ps[s2.pc], "cats")

    def test_format_3_reversed_flat(self):
        s = Session()
        s.load_table(ANIMALS_TABLE, id_col=0, parent_cols=[1], fmt=0)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "flat.csv")
            self.assertTrue(s.export_flat(path, hier="Parent", reverse=True, overwrite=True)["ok"])
            with open(path, newline="", encoding="utf-8") as fh:
                rows = list(csv.reader(fh))
        hier_cols = [i for i, name in enumerate(rows[0]) if name.startswith("Parent_")]
        s2 = Session()
        out = s2.load_table(rows, parent_cols=hier_cols, fmt=3)
        self.assertTrue(out["ok"], out)
        self.assertEqual(sorted(s2.nodes), ["animals", "cats", "lion"])
        self.assertEqual(s2.nodes["lion"].ps[s2.pc], "cats")

    def test_format_5_indented(self):
        s = Session()
        out = s.load_table(FMT5, fmt=5)
        self.assertTrue(out["ok"], out)
        self.assertEqual(sorted(s.nodes), ["animals", "cats", "lion"])
        self.assertEqual(s.nodes["cats"].ps[s.pc], "animals")

    def test_format_6_indented_multi_detail(self):
        s = Session()
        out = s.load_table(FMT6, fmt=6)
        self.assertTrue(out["ok"], out)
        self.assertEqual(sorted(s.nodes), ["animals", "cats", "lion"])
        self.assertGreaterEqual(len(s.headers), 3)

    def test_format_7_indented_header(self):
        s = Session()
        out = s.load_table(FMT7, fmt=7)
        self.assertTrue(out["ok"], out)
        self.assertEqual(sorted(s.nodes), ["animals", "cats", "lion"])
        self.assertEqual(s.nodes["lion"].ps[s.pc], "cats")

    def test_format_1_needs_parents(self):
        s = Session()
        out = s.load_table([["A", "B"], ["x", "y"]], fmt=1)
        self.assertFalse(out["ok"])
        self.assertEqual(out["error"]["code"], "need_columns")


class TestXlsxSheetFlags(unittest.TestCase):
    def test_need_sheet_when_workbook_has_two_sheets(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "two.xlsx")
            wb = Workbook()
            wb.active.title = "One"
            wb.active.append(["ID", "Parent", "Name"])
            wb.active.append(["Animals", "", "All"])
            ws2 = wb.create_sheet("Two")
            ws2.append(["ID", "Parent"])
            wb.save(path)
            out = Session().load_path(path)
            self.assertEqual(out["error"]["code"], "need_sheet")

    def test_need_sheet_named_and_index(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "two.xlsx")
            wb = Workbook()
            wb.active.title = "One"
            wb.active.append(["ID", "Parent", "Name"])
            wb.active.append(["Animals", "", "All"])
            wb.create_sheet("Two").append(["ID", "Parent"])
            wb.save(path)
            named = Session().load_path(path, sheet="One", id_col=0, parent_cols=[1])
            self.assertTrue(named["ok"], named)
            indexed = Session().load_path(path, sheet=0, id_col=0, parent_cols=[1])
            self.assertTrue(indexed["ok"], indexed)
            missing = Session().load_path(path, sheet="Nope", id_col=0, parent_cols=[1])
            self.assertEqual(missing["error"]["code"], "need_sheet")

    def test_program_data_present_rejects_column_flags(self):
        s = Session()
        s.load_table(ANIMALS_TABLE, id_col=0, parent_cols=[1], fmt=0)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "app.xlsx")
            self.assertTrue(s.save_path(path, overwrite=True)["ok"])
            flagged = Session().load_path(path, id_col=0, parent_cols=[1])
            self.assertEqual(flagged["error"]["code"], "program_data_present")
            s2 = Session()
            self.assertTrue(s2.load_path(path)["ok"])
            self.assertEqual(sorted(s2.nodes), ["animals", "cats", "lion"])
