# SPDX-License-Identifier: AGPL-3.0-only
# Copyright (c) R. A. Gardner

from __future__ import annotations

import csv
import datetime
import json
import os
import pickle
import re
import zlib
from bisect import bisect_left
from collections import defaultdict, deque
from itertools import filterfalse, repeat
from operator import itemgetter

from openpyxl import Workbook, load_workbook

from .changelog import (
    ORIGIN_IMPORT,
    ORIGIN_MERGE,
    ChangeBuilder,
    flatten_changelog,
    is_member_type,
    load_changelog,
    strip_type,
)
from .classes import Header, Node, RowStorage, TreeBuilder, normalize_header_type
from .constants import software_version_number
from .functions import (
    b32_x_dict,
    bytes_io_wb,
    case_insensitive_replace,
    csv_str_x_data,
    dict_x_b32,
    equalize_sublist_lens,
    full_sheet_to_dict,
    get_json_format,
    get_json_from_file,
    json_to_sheet,
    sort_key,
    try_remove,
    ws_x_data,
    ws_x_program_data_str,
    xl_column_string,
)

_DATA_SUFFIXES = (".xlsx", ".xls", ".xlsm", ".csv", ".tsv", ".json")


def ok(result=None, warnings=None) -> dict:
    return {
        "ok": True,
        "error": None,
        "warnings": list(warnings or []),
        "result": result if result is not None else {},
    }


def fail(code: str, message: str, warnings=None, **extra) -> dict:
    err = {"code": code, "message": message, **extra}
    return {
        "ok": False,
        "error": err,
        "warnings": list(warnings or []),
        "result": {},
    }


def _alpha_to_idx(token: str):
    if not token or not str(token).isalpha():
        return None
    n = 0
    for c in str(token).upper():
        n = n * 26 + ord(c) - 64
    return n - 1


def _move_elements(seq, new_idxs):
    old_idxs = dict(zip(new_idxs.values(), new_idxs))
    remaining_values = (e for i, e in enumerate(seq) if i not in new_idxs)
    return [seq[old_idxs[i]] if i in old_idxs else next(remaining_values) for i in range(len(seq))]


def _full_new_idxs(n, new_idxs):
    moved = _move_elements(list(range(n)), new_idxs)
    return {old: new for new, old in enumerate(moved)}


class ListOps:
    def __init__(self, session: Session):
        self.s = session

    def insert_rows(self, rows, idx=None):
        if idx is None:
            self.s.data.extend(rows)
        else:
            for i, row in enumerate(rows):
                self.s.data.insert(idx + i, row)

    def delete_rows(self, idxs):
        for i in sorted(idxs, reverse=True):
            del self.s.data[i]

    def insert_cols(self, idx, n=1):
        for row in self.s.data:
            row[idx:idx] = [""] * n

    def delete_cols(self, idxs):
        drop = sorted(idxs, reverse=True)
        for row in self.s.data:
            for i in drop:
                del row[i]


class Session:
    def __init__(self, ops=None):
        self.data: list[list[str]] = []
        self.ops = ops or ListOps(self)
        self.nodes: dict = {}
        self.rns: dict[str, int] = {}
        self.headers: list = []
        self.ic = 0
        self.pc = 0
        self.hiers: list[int] = []
        self.row_len = 0
        self.changelog: list = []
        self._pending = None
        self.warnings: list = []
        self.tagged_ids: set = set()
        self.topnodes_order: dict = {}
        self.auto_sort_nodes_bool = True
        self.allow_spaces_ids_var = False
        self.allow_spaces_columns_var = False
        self.vs = deque(maxlen=30)
        self.unsaved = False
        self.filepath = None
        self.sheetname = "Sheet1"
        self.json_format = 1
        self.save_xlsx_with_program_data = True
        self.save_json_with_program_data = True
        self.refresh_rows = set()
        self.sort_later_dct = None
        self.changelog_at_open = 0
        self.has_document = False
        self.on_increment_unsaved = lambda: None
        self.on_changelog_row = lambda: None

    def get_datetime_changelog(self, increment_unsaved=True):
        if increment_unsaved:
            self.unsaved = True
            self.on_increment_unsaved()
        return f"{datetime.datetime.today().strftime('%Y/%m/%d')}"

    def commit_change(self, builder, increment_unsaved=True):
        at = self.get_datetime_changelog(increment_unsaved=increment_unsaved)
        ch = builder.finish(at)
        self.changelog.append(ch)
        self._pending = None
        self.on_changelog_row()
        return ch

    def _log(self, kind, what="", old="", new="", increment_unsaved=True):
        b = ChangeBuilder()
        b.add(kind, what, old, new)
        return self.commit_change(b, increment_unsaved=increment_unsaved)

    def _log_prefixed(self, builder, typ, what="", old="", new=""):
        builder.add(strip_type(typ)[1], what, old, new)

    def _drop_pending(self):
        self._pending = None

    def _commit_pending(self, increment_unsaved=True):
        b = self._pending
        self._pending = None
        if b is None or not b.rows:
            return None
        return self.commit_change(b, increment_unsaved=increment_unsaved)

    def _pending_add(self, change, id_, old, new):
        origin, kind = strip_type(change)
        if self._pending is None:
            self._pending = ChangeBuilder(origin=origin)
        elif self._pending.origin != origin:
            self._commit_pending()
            self._pending = ChangeBuilder(origin=origin)
        self._pending.add(kind, id_, old, new)

    def changelog_append(self, change, id_, old, new):
        # Shim: GUI paste/import/merge still emit member rows then a summary.
        # Produces one Change. Prefer ChangeBuilder + commit_change for new code.
        if self._pending is not None:
            if is_member_type(change):
                self._pending_add(change, id_, old, new)
                return
            self._pending.summary(change, id_, old, new)
            self._commit_pending()
            return
        if is_member_type(change):
            self._pending_add(change, id_, old, new)
            return
        self._log(change, id_, old, new)

    def changelog_append_no_unsaved(self, change, id_, old, new):
        self._pending_add(change, id_, old, new)

    def changelog_singular(self, text):
        if self._pending is None:
            if self.changelog:
                last = self.changelog[-1]
                last.label = text
                if last.rows:
                    last.rows[0].kind = text
                    last.rows[0].display_type = text
            self.unsaved = True
            self.on_increment_unsaved()
            return
        self._pending.label = text
        self._pending.force_singular = True
        self._commit_pending()

    def clear(self):
        self.data = []
        self.nodes = {}
        self.rns = {}
        self.headers = []
        self.ic = 0
        self.pc = 0
        self.hiers = []
        self.row_len = 0
        self.changelog = []
        self._pending = None
        self.warnings = []
        self.tagged_ids = set()
        self.topnodes_order = {}
        self.auto_sort_nodes_bool = True
        self.vs = deque(maxlen=30)
        self.unsaved = False
        self.filepath = None
        self.refresh_rows = set()
        self.sort_later_dct = None
        self.changelog_at_open = 0
        self.has_document = False

    def status_dict(self):
        hierarchy = None
        if self.headers and 0 <= self.pc < len(self.headers):
            hierarchy = self.headers[self.pc].name
        return {
            "file": None if self.filepath is None else self.filepath,
            "unsaved": self.unsaved,
            "hierarchy": hierarchy,
            "ids": len(self.nodes),
            "undo": len(self.vs),
        }

    def _column_entry(self, index):
        header = self.headers[index]
        if index == self.ic:
            role = "id"
        elif index in self.hiers:
            role = "parent"
        else:
            role = "detail"
        return {
            "index": index,
            "letter": xl_column_string(index + 1),
            "name": header.name,
            "role": role,
            "type": header.type_,
        }

    def _open_result(self, fmt):
        columns = [self._column_entry(i) for i in range(len(self.headers))]
        id_column = self._column_entry(self.ic) if self.headers else None
        hierarchies = [
            {
                "index": h,
                "letter": xl_column_string(h + 1),
                "name": self.headers[h].name,
            }
            for h in self.hiers
            if h < len(self.headers)
        ]
        return {
            "ids": len(self.nodes),
            "format": fmt,
            "id_column": id_column,
            "columns": columns,
            "hierarchies": hierarchies,
            "warnings": list(self.warnings),
        }

    def fix_headers(self, headers, row_len, warnings=True):
        if len(headers) < row_len:
            headers += list(repeat("", row_len - len(headers)))
        tally_of_headers = defaultdict(lambda: -1)
        allow_whitespace = self.allow_spaces_columns_var
        for coln in range(len(headers)):
            cell = headers[coln]
            if not cell:
                cell = f"MISSING_{coln + 1}"
                if warnings:
                    self.warnings.append(f" - Missing header in column #{coln + 1}")
            if not allow_whitespace:
                if warnings:
                    if " " in cell:
                        self.warnings.append(f" - Spaces in header column #{coln + 1}")
                    if "\n" in cell:
                        self.warnings.append(f" - Newlines in header column #{coln + 1}")
                    if "\r" in cell:
                        self.warnings.append(f" - Carriage returns in header column #{coln + 1}")
                    if "\t" in cell:
                        self.warnings.append(f" - Tabs in header column #{coln + 1}")
                cell = "".join(cell.strip().split())
            hk = cell.lower()
            tally_of_headers[hk] += 1
            if tally_of_headers[hk] > 0:
                if warnings:
                    self.warnings.append(f" - Duplicate header in column #{coln + 1}")
                orig = cell
                x = 1
                while hk in tally_of_headers:
                    cell = f"{orig}_DUPLICATED_{x}"
                    hk = cell.lower()
                    x += 1
                tally_of_headers[hk] += 1
            headers[coln] = cell
        return headers

    def sort_node_cn(self, cn: list[str], h: int):
        wc = []
        woc = []
        for ciid in cn:
            if self.nodes[ciid].cn[h]:
                wc.append(ciid)
            else:
                woc.append(ciid)
        return sorted(wc, key=sort_key) + sorted(woc, key=sort_key)

    def remake_topnodes_order(self):
        self.topnodes_order = {}
        for h in self.hiers:
            wc = []
            woc = []
            for iid, node in self.nodes.items():
                if node.ps[h] == "":
                    if node.cn[h]:
                        wc.append(iid)
                    else:
                        woc.append(iid)
            self.topnodes_order[h] = sorted(wc, key=sort_key) + sorted(woc, key=sort_key)

    def associate(self, startup=True):
        first_hier = self.hiers[0]
        quick_hiers = self.hiers[1:]
        lh = len(self.hiers)
        if startup and self.auto_sort_nodes_bool:
            for node in self.nodes.values():
                if all(p is None for p in node.ps.values()):
                    node.ps = {h: "" if node.cn[h] else None for h in self.hiers}
                    newrow = list(repeat("", self.row_len))
                    newrow[self.ic] = node.name
                    self.data.append(newrow)
                    self.warnings.append(f" - ID ({node.name}) missing from ID column, new row added")
                tlly = 0
                for k, v in node.cn.items():
                    if v:
                        node.cn[k] = self.sort_node_cn(v, k)
                    elif not node.ps[k]:
                        node.ps[k] = None
                        tlly += 1
                if tlly == lh:
                    node.ps[first_hier] = ""
                    for h in quick_hiers:
                        node.ps[h] = None

        elif not startup and self.auto_sort_nodes_bool:
            to_insert = []
            for node in self.nodes.values():
                if all(p is None for p in node.ps.values()):
                    node.ps = {h: "" if node.cn[h] else None for h in self.hiers}
                    newrow = list(repeat("", self.row_len))
                    newrow[self.ic] = node.name
                    to_insert.append(newrow)
                tlly = 0
                for k, v in node.cn.items():
                    if v:
                        node.cn[k] = self.sort_node_cn(v, k)
                    elif not node.ps[k]:
                        node.ps[k] = None
                        tlly += 1
                if tlly == lh:
                    node.ps[first_hier] = ""
                    for h in quick_hiers:
                        node.ps[h] = None
            if to_insert:
                self.ops.insert_rows(to_insert)

        elif not startup and not self.auto_sort_nodes_bool:
            st_check_topnodes_order = {k: set(v) for k, v in self.topnodes_order.items()}
            to_insert = []
            for iid, node in self.nodes.items():
                if all(p is None for p in node.ps.values()):
                    node.ps = {h: "" if node.cn[h] else None for h in self.hiers}
                    newrow = list(repeat("", self.row_len))
                    newrow[self.ic] = node.name
                    to_insert.append(newrow)
                tlly = 0
                for k, v in node.cn.items():
                    if not v and not node.ps[k]:
                        node.ps[k] = None
                        tlly += 1
                if tlly == lh:
                    if all(iid not in h for h in st_check_topnodes_order.values()):
                        node.ps[first_hier] = ""
                        for h in quick_hiers:
                            node.ps[h] = None
                        self.topnodes_order[first_hier].append(iid)
                    else:
                        for h, v in st_check_topnodes_order.items():
                            if iid in v:
                                node.ps[h] = ""
            if to_insert:
                self.ops.insert_rows(to_insert)

    def nodes_json_x_dict(self, njson: dict, hiers) -> dict:
        return {
            name.lower(): Node(
                name=name,
                hrs=hiers,
                cn={int(h): cnl for h, cnl in nodedict["cn"].items()},
                ps={int(h): pk for h, pk in nodedict["ps"].items()},
            )
            for name, nodedict in njson.items()
        }

    def jsonify_nodes(self):
        return {
            n.name: {
                "cn": n.cn,
                "ps": n.ps,
            }
            for n in self.nodes.values()
        }

    def gen_sheet_w_headers(self):
        yield (h.name for h in self.headers)
        yield from ((e if e else None for e in r) for r in self.data)

    def program_data_dict(self, sheetname="n/a"):
        return {
            "records": self.data,
            "ic": self.ic,
            "pc": self.pc,
            "hiers": self.hiers,
            "headers": [
                {
                    "name": h.name,
                    "type": h.type_,
                    "formatting": h.formatting,
                    "validation": h.validation,
                }
                for h in self.headers
            ],
            "nodes": self.jsonify_nodes(),
            "changelog_format": 2,
            "changelog": [ch.to_dict() for ch in self.changelog],
            "topnodes_order": self.topnodes_order,
            "tagged_ids": list(self.tagged_ids),
            "auto_sort_nodes_bool": self.auto_sort_nodes_bool,
            "sheetname": sheetname,
            "allow_spaces_ids": self.allow_spaces_ids_var,
            "allow_spaces_columns": self.allow_spaces_columns_var,
        }

    def load_table(self, rows, *, id_col=None, parent_cols=None, fmt=0, associate=True, strip_ids=True):
        rows = [list(r) for r in rows]
        if not rows:
            return fail("invalid_format", "File contains no data   ")
        parent_cols = list(parent_cols) if parent_cols is not None else []
        if fmt in (1, 2, 3, 4) and not parent_cols:
            return fail("need_columns", "Need parent columns   ")
        if fmt not in (1, 2, 3, 4, 5, 6, 7) and (id_col is None or not parent_cols or id_col in parent_cols):
            return fail("need_columns", "Need ID and parent columns   ")
        self.warnings = []
        self.nodes = {}
        row_len = equalize_sublist_lens(rows)
        if fmt in (1, 2, 3, 4):
            rows, row_len, id_col, parent_cols = TreeBuilder().convert_flattened_to_normal(
                data=rows,
                hier_cols=parent_cols,
                rowlen=row_len,
                fmt=fmt,
                warnings=self.warnings,
            )
        elif fmt == 5:
            rows, row_len, id_col, parent_cols = TreeBuilder().convert_indented_tree_detail_adjacent_to_normal(
                data=rows,
            )
        elif fmt == 6:
            rows, row_len, id_col, parent_cols = TreeBuilder().convert_indented_tree_details_adjacent_to_normal(
                data=rows,
            )
        elif fmt == 7:
            rows, row_len, id_col, parent_cols = TreeBuilder().convert_indented_tree_with_header_to_normal(
                data=rows,
            )
        if not rows:
            return fail("invalid_format", "File contains no data   ")
        names = self.fix_headers(rows.pop(0), row_len)
        self.headers = [
            Header(
                name,
                type_="ID" if i == id_col else "Parent" if i in parent_cols else "Text",
            )
            for i, name in enumerate(names)
        ]
        self.ic = int(id_col)
        self.hiers = list(parent_cols)
        self.pc = int(self.hiers[0])
        self.row_len = int(row_len)
        built, nodes, warnings = TreeBuilder().build(
            input_sheet=rows,
            output_sheet=[],
            row_len=self.row_len,
            ic=self.ic,
            hiers=self.hiers,
            nodes={},
            warnings=self.warnings,
            strip=strip_ids,
            fix_associate=False,
        )
        self.data = built
        self.nodes = nodes
        self.warnings = warnings
        if associate:
            self.associate(startup=True)
        self.rns = {row[self.ic].lower(): i for i, row in enumerate(self.data)}
        self.remake_topnodes_order()
        self.tagged_ids = set()
        self.has_document = True
        return ok(self._open_result(fmt), warnings=list(self.warnings))

    def load_program_data(self, d):
        self.data = d["records"]
        self.ic = int(d["ic"])
        self.pc = int(d["pc"])
        self.hiers = [int(h) for h in d["hiers"]]
        self.headers = [
            Header(
                h["name"],
                h["type"],
                [tuple(x) for x in h["formatting"]],
                h["validation"],
            )
            for h in d["headers"]
        ]
        self.row_len = len(self.headers)
        self._pending = None
        self.changelog = load_changelog(d.get("changelog") or [], warnings=self.warnings)
        self.allow_spaces_ids_var = bool(d["allow_spaces_ids"])
        self.allow_spaces_columns_var = bool(d["allow_spaces_columns"])
        self.auto_sort_nodes_bool = bool(d["auto_sort_nodes_bool"])
        try:
            self.nodes = self.nodes_json_x_dict(d["nodes"], hiers=self.hiers)
        except Exception:
            self.warnings = []
            built, nodes, warnings = TreeBuilder().build(
                input_sheet=list(self.data),
                output_sheet=[],
                row_len=self.row_len,
                ic=self.ic,
                hiers=self.hiers,
                nodes={},
                warnings=[],
                strip=not self.allow_spaces_ids_var,
                fix_associate=False,
            )
            self.data = built
            self.nodes = nodes
            self.warnings = warnings
        self.topnodes_order = {int(h): v for h, v in d["topnodes_order"].items()}
        self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
        self.tagged_ids = set(d["tagged_ids"])
        self.has_document = True
        self.unsaved = False
        self.changelog_at_open = len(self.changelog)
        self.vs = deque(maxlen=30)
        return ok(self._open_result(0), warnings=list(self.warnings))

    def new(self, *, discard=False):
        if self.unsaved and not discard:
            return fail("unsaved", "Unsaved changes")
        self.data = []
        self.headers = [
            Header("ID", "ID"),
            Header("DETAIL_1"),
            Header("PARENT_1", "Parent"),
        ]
        self.ic = 0
        self.pc = 2
        self.hiers = [2]
        self.row_len = 3
        self.nodes = {}
        self.rns = {}
        self.changelog = []
        self._pending = None
        self.warnings = []
        self.tagged_ids = set()
        self.topnodes_order = {}
        self.auto_sort_nodes_bool = True
        self.vs = deque(maxlen=30)
        self.filepath = None
        self.sheetname = "Sheet1"
        self.has_document = True
        self.changelog_at_open = 0
        self.unsaved = False
        self.refresh_rows = set()
        self.sort_later_dct = None
        self.remake_topnodes_order()
        return ok(self._open_result(0), warnings=[])

    def close(self, *, discard=False):
        if self.unsaved and not discard:
            return fail("unsaved", "Unsaved changes")
        self.clear()
        return ok({})

    def _load_rows_or_need_columns(self, rows, *, id_col, parent_cols, fmt):
        if not rows:
            return fail("invalid_format", "File contains no data   ")
        if fmt in (5, 6, 7):
            return self.load_table(rows, fmt=fmt)
        if fmt in (1, 2, 3, 4):
            if not parent_cols:
                return fail("need_columns", "Need parent columns   ")
            return self.load_table(rows, parent_cols=parent_cols, fmt=fmt)
        if id_col is None or not parent_cols:
            return fail("need_columns", "Need ID and parent columns   ")
        return self.load_table(rows, id_col=id_col, parent_cols=parent_cols, fmt=fmt)

    def load_path(self, path, *, sheet=None, id_col=None, parent_cols=None, fmt=0, discard=False):
        if self.unsaved and not discard:
            return fail("unsaved", "Unsaved changes")
        if not os.path.isfile(path):
            return fail("file_not_found", f"File not found: {path}")
        suffix = os.path.splitext(path)[1].lower()
        if suffix not in _DATA_SUFFIXES:
            return fail("invalid_format", "File must be .xlsx, .xlsm, .xls, .csv, .tsv or .json   ")
        column_flags = id_col is not None or parent_cols is not None or fmt != 0
        parent_cols = list(parent_cols) if parent_cols is not None else None

        if suffix in (".csv", ".tsv"):
            try:
                with open(path, "r") as fh:
                    rows = csv_str_x_data(fh.read())
            except Exception as error_msg:
                return fail("invalid_format", f"{error_msg}")
            out = self._load_rows_or_need_columns(rows, id_col=id_col, parent_cols=parent_cols, fmt=fmt)
            if not out["ok"]:
                return out
            self.sheetname = "Sheet1"

        elif suffix == ".json":
            try:
                j = get_json_from_file(path)
            except Exception as error_msg:
                return fail("invalid_format", f"{error_msg}")
            if "program_data" in j:
                if column_flags:
                    return fail("program_data_present", "File already has app data   ")
                try:
                    program_data = b32_x_dict(j["program_data"])
                    out = self.load_program_data(program_data)
                except Exception as error_msg:
                    return fail("invalid_format", f"{error_msg}")
                if not out["ok"]:
                    return out
                self.sheetname = "Sheet1"
            else:
                json_format = get_json_format(j)
                if not json_format:
                    return fail("invalid_format", "Could not find data of correct format   ")
                rows, _row_len = json_to_sheet(
                    j,
                    format_=json_format[0],
                    key=json_format[1],
                    get_format=False,
                    return_rowlen=True,
                )
                out = self._load_rows_or_need_columns(rows, id_col=id_col, parent_cols=parent_cols, fmt=fmt)
                if not out["ok"]:
                    return out
                self.sheetname = "Sheet1"

        else:
            try:
                in_mem = bytes_io_wb(path)
                wb = load_workbook(in_mem, read_only=True, data_only=True)
            except Exception as error_msg:
                return fail("invalid_format", f"{error_msg}")
            try:
                if len(wb.sheetnames) < 1:
                    return fail("invalid_format", "File contains no data   ")
                if "program_data" in wb.sheetnames:
                    if column_flags:
                        return fail("program_data_present", "File already has app data   ")
                    ws = wb["program_data"]
                    ws.reset_dimensions()
                    try:
                        program_data = b32_x_dict(ws_x_program_data_str(ws))
                        out = self.load_program_data(program_data)
                    except Exception as error_msg:
                        return fail("invalid_format", f"{error_msg}")
                    if not out["ok"]:
                        return out
                    self.sheetname = program_data["sheetname"]
                else:
                    names = wb.sheetnames
                    if sheet is None:
                        if len(names) > 1:
                            return fail("need_sheet", "Workbook has more than one sheet   ")
                        chosen = names[0]
                    elif isinstance(sheet, int):
                        if sheet < 0 or sheet >= len(names):
                            return fail("need_sheet", "Sheet index not found   ")
                        chosen = names[sheet]
                    else:
                        try:
                            wb[sheet]
                            chosen = sheet
                        except Exception:
                            return fail("need_sheet", f"Sheet not found: {sheet}")
                    ws = wb[chosen]
                    ws.reset_dimensions()
                    rows = ws_x_data(ws)
                    out = self._load_rows_or_need_columns(rows, id_col=id_col, parent_cols=parent_cols, fmt=fmt)
                    if not out["ok"]:
                        return out
                    self.sheetname = chosen
            finally:
                wb.close()

        self.filepath = os.path.abspath(path)
        self.unsaved = False
        self.has_document = True
        self.changelog_at_open = len(self.changelog)
        result = dict(out["result"])
        result["warnings"] = list(self.warnings)
        return ok(result, warnings=list(self.warnings))

    def _xlsx_chunker(self, seq):
        size = min(len(seq), 32000)
        return (seq[pos : pos + size] for pos in range(0, len(seq), size))

    def save_path(self, path, *, overwrite=False, sheetname=None):
        if os.path.exists(path) and not overwrite:
            return fail("file_exists", "File already exists   ")
        suffix = os.path.splitext(path)[1].lower()
        if suffix in (".xls", ".xlsm") or suffix not in (".csv", ".tsv", ".json", ".xlsx"):
            return fail("invalid_format", "Can only write .json, .xlsx or .csv    ")
        name = sheetname if sheetname is not None else self.sheetname
        if suffix in (".csv", ".tsv"):
            kind = "tsv" if suffix == ".tsv" else "csv"
            with open(path, "w", newline="", encoding="utf-8") as fh:
                writer = csv.writer(
                    fh,
                    dialect=csv.excel_tab if kind == "tsv" else csv.excel,
                    lineterminator="\n",
                )
                writer.writerows(self.gen_sheet_w_headers())
        elif suffix == ".json":
            kind = "json"
            d = full_sheet_to_dict(
                [h.name for h in self.headers],
                self.data,
                format_=self.json_format,
            )
            if self.save_json_with_program_data:
                d["version"] = software_version_number
                d["changelog"] = flatten_changelog(self.changelog)[0]
                d["program_data"] = dict_x_b32(self.program_data_dict(name))
            with open(path, "w") as fh:
                fh.write(json.dumps(d, indent=4))
        else:
            kind = "xlsx"
            wb = Workbook(write_only=True)
            ws = wb.create_sheet(title=name)
            if not self.ic:
                ws.freeze_panes = "B2"
            else:
                ws.freeze_panes = "A2"
            for row in self.gen_sheet_w_headers():
                ws.append(row)
            if self.save_xlsx_with_program_data:
                hidden = wb.create_sheet(title="program_data")
                hidden.append([f"{software_version_number}"])
                for chunk in self._xlsx_chunker(dict_x_b32(self.program_data_dict(name))):
                    hidden.append([chunk])
                hidden.sheet_state = "hidden"
            wb.active = wb[name]
            wb.save(path)
        self.filepath = os.path.abspath(path)
        self.sheetname = name
        self.unsaved = False
        return ok({"file": self.filepath, "kind": kind})

    def check_cn(self, iid: str, h: int):
        stack = [iid]
        while stack:
            current = stack.pop()
            yield current
            stack.extend(reversed(self.nodes[current].cn[h]))

    def check_ps(self, iid: str, h: int):
        current = iid
        while True:
            yield current
            if not self.nodes[current].ps[h]:
                break
            current = self.nodes[current].ps[h]

    def get_ids_parent(self, iid) -> str:
        if self.nodes[iid.lower()].ps[self.pc]:
            return self.nodes[iid.lower()].ps[self.pc]
        return ""

    def _untag_id(self, ik):
        self.tagged_ids.discard(ik)

    def is_in_validation(self, validation, text):
        return text in validation

    def detail_is_valid_for_col(self, col, detail):
        return not (self.headers[col].validation and not self.is_in_validation(self.headers[col].validation, detail))

    def _copy_headers(self):
        return [
            Header(
                f"{h.name}",
                f"{h.type_}",
                [tuple(t) for t in h.formatting],
                h.validation.copy(),
            )
            for h in self.headers
        ]

    def snapshot_required_data(self):
        return {
            "topnodes_order": {k: list(v) for k, v in self.topnodes_order.items()},
            "tagged_ids": set(self.tagged_ids),
            "headers": self._copy_headers(),
            "ic": int(self.ic),
            "pc": int(self.pc),
            "hiers": list(self.hiers),
            "row_len": int(self.row_len),
            "auto_sort_nodes_bool": bool(self.auto_sort_nodes_bool),
            "nodes": pickle.dumps(self.nodes),
        }

    def snapshot_add_id(self):
        self.vs.append(
            {
                "type": "add id",
                "row": {},
                "required_data": self.snapshot_required_data(),
            }
        )

    def snapshot_rename_id(self):
        self.vs.append(
            {
                "type": "rename id",
                "rows": [],
                "ikrow": (),
                "required_data": self.snapshot_required_data(),
            }
        )

    def snapshot_delete_ids(self):
        self.vs.append(
            {
                "type": "delete ids",
                "rows": {},
                "required_data": self.snapshot_required_data(),
            }
        )

    def _ensure_snapshot(self, type_, maker):
        if not self.vs or self.vs[-1].get("type") != type_:
            maker()

    def _id_has_spaces(self, name: str) -> bool:
        return bool(re.search(r"[\n\t\s]", name))

    def _splice_tree_order(self, ik, sibling, before):
        sk = sibling.lower()
        parent_key = self.nodes[ik].ps[self.pc]
        order = self.topnodes_order[self.pc] if parent_key == "" else self.nodes[parent_key].cn[self.pc]
        if ik in order:
            order.remove(ik)
        at = order.index(sk)
        order.insert(at if before else at + 1, ik)

    def add(self, ID, parent, insert_row=None, snapshot=True, *, before=None, after=None):
        if not ID:
            return fail("empty_id", "New name cannot be empty   ")
        if not self.allow_spaces_ids_var and self._id_has_spaces(ID):
            return fail("spaces_not_allowed", "Spaces are not allowed in IDs   ")
        if before is not None and after is not None:
            return fail("usage", "Use only one of before or after   ")
        if before is not None or after is not None:
            if self.auto_sort_nodes_bool:
                return fail("auto_sort_on", "Turn auto-sort off to use before/after   ")
            sib = before if before is not None else after
            sk = sib.lower()
            if sk not in self.nodes:
                return fail("id_not_found", "ID doesn't exist   ", id=sib)
            if self.nodes[sk].ps[self.pc] is None:
                return fail("not_in_hierarchy", f"{sib} is not in this hierarchy   ")
            implied_key = self.nodes[sk].ps[self.pc]
            implied = "" if implied_key == "" else self.nodes[implied_key].name
            if parent and implied.lower() != parent.lower():
                return fail("usage", "before/after cannot be combined with parent   ")
            parent = implied
            order = self.topnodes_order[self.pc] if implied_key == "" else self.nodes[implied_key].cn[self.pc]
            if sk not in order:
                return fail("not_sibling", f"{sib} is not a sibling in this hierarchy   ")
        pk = parent.lower()
        if parent:
            if pk not in self.nodes:
                return fail("id_not_found", "ID doesn't exist   ", id=parent)
            if self.nodes[pk].ps[self.pc] is None:
                return fail("not_in_hierarchy", f"{parent} is not in this hierarchy   ")
        ik = ID.lower()
        if ik in self.nodes and self.nodes[ik].ps[self.pc] is not None:
            return fail("already_in_hierarchy", "ID already in hierarchy   ", id=ID)
        created_row = ik not in self.nodes
        if snapshot:
            self._ensure_snapshot("add id", self.snapshot_add_id)
        if ik not in self.nodes:
            self.nodes[ik] = Node(ID, self.hiers)
            newrow = list(repeat("", self.row_len))
            newrow[self.ic] = ID
            newrow[self.pc] = parent
            self.ops.insert_rows([newrow], insert_row)
            rn = len(self.data) - 1 if insert_row is None else int(insert_row)
            self.rns[ik] = rn
            if snapshot:
                self.vs[-1]["row"]["added_or_changed"] = "added"
                self.vs[-1]["row"]["rn"] = rn
        else:
            rn = self.rns[ik]
            if snapshot:
                self.vs[-1]["row"]["added_or_changed"] = "changed"
                self.vs[-1]["row"]["rn"] = rn
                self.vs[-1]["row"]["stored"] = self.data[rn].copy()
            self.data[rn][self.pc] = parent
        if parent == "":
            self.nodes[ik].ps[self.pc] = ""
        else:
            self.nodes[ik].ps[self.pc] = pk
            self.nodes[pk].cn[self.pc].append(ik)
            if self.auto_sort_nodes_bool:
                self.nodes[pk].cn[self.pc] = self.sort_node_cn(self.nodes[pk].cn[self.pc], self.pc)
                if self.nodes[pk].ps[self.pc]:
                    parent_parent_node = self.nodes[self.nodes[pk].ps[self.pc]]
                    parent_parent_node.cn[self.pc] = self.sort_node_cn(parent_parent_node.cn[self.pc], self.pc)
        if not self.auto_sort_nodes_bool and parent == "":
            self.topnodes_order[self.pc].append(ik)
        if insert_row is not None and snapshot:
            self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
        if before is not None or after is not None:
            self._splice_tree_order(ik, before if before is not None else after, before is not None)
        if snapshot:
            parent_text = parent if parent else "n/a - Top ID"
            self.changelog_append(
                "Add ID",
                f"Name: {ID} Parent: {parent_text} column #{self.pc + 1} named: {self.headers[self.pc].name}",
                "",
                "",
            )
        display = self.nodes[ik].name
        return ok(
            {
                "added": [
                    {
                        "id": display,
                        "parent": parent,
                        "hierarchy": self.headers[self.pc].name,
                        "created_row": created_row,
                    }
                ]
            }
        )

    def rename(self, ID, new_name, snapshot=True):
        ik = ID.lower()
        if ik not in self.nodes:
            return fail("id_not_found", "ID doesn't exist   ")
        nnk = new_name.lower()
        if nnk in self.nodes and ik != nnk:
            return fail("id_exists", "New name already exists   ")
        if not nnk:
            return fail("empty_id", "New name cannot be empty   ")
        if not self.allow_spaces_ids_var and self._id_has_spaces(new_name):
            return fail("spaces_not_allowed", "Spaces are not allowed in IDs   ")
        self.refresh_rows = set()
        if snapshot:
            self._ensure_snapshot("rename id", self.snapshot_rename_id)
            qvsrwsapp = self.vs[-1]["rows"].append
        ik_rn = self.rns[ik]
        self.data[ik_rn][self.ic] = new_name
        for h, cn in self.nodes[ik].cn.items():
            for ciid in cn:
                chld_rn = self.rns[ciid]
                self.refresh_rows.add(chld_rn)
                if snapshot:
                    qvsrwsapp(zlib.compress(pickle.dumps((chld_rn, h, self.data[chld_rn][h]))))
                self.data[chld_rn][h] = f"{new_name}"
                self.nodes[ciid].ps[h] = nnk
        for h, p in self.nodes[ik].ps.items():
            if p:
                self.nodes[p].cn[h][self.nodes[p].cn[h].index(ik)] = nnk
        if snapshot:
            self.vs[-1]["ikrow"] = (ik_rn, ik, self.nodes[ik].name, new_name)
        old = self.nodes[ik].name
        self.nodes[ik].name = new_name
        self.nodes[nnk] = self.nodes.pop(ik)
        self.rns[nnk] = self.rns.pop(ik)
        if self.auto_sort_nodes_bool:
            for h, p in self.nodes[nnk].ps.items():
                if p:
                    parent_node = self.nodes[self.nodes[nnk].ps[h]]
                    parent_node.cn[h] = self.sort_node_cn(parent_node.cn[h], h)
        else:
            for h in self.hiers:
                if self.nodes[nnk].ps[h] == "":
                    try:
                        self.topnodes_order[h][self.topnodes_order[h].index(ik)] = nnk
                    except Exception:
                        continue
        if ik in self.tagged_ids:
            self.tagged_ids.discard(ik)
            self.tagged_ids.add(nnk)
        if snapshot:
            self.changelog_append("Rename ID", old, old, new_name)
        return ok({"old": old, "new": new_name})

    def set_detail(self, ID, col, value, snapshot=True):
        ik = ID.lower()
        if ik not in self.nodes:
            return fail("id_not_found", "ID doesn't exist   ")
        if col == self.ic or col in self.hiers:
            return fail("usage", "Use rename for the ID column and move for parent columns   ")
        if not (0 <= col < len(self.headers)):
            return fail("unknown_column", "Unknown column   ")
        if not self.detail_is_valid_for_col(col, value):
            return fail("validation", "Value is not in the column validation list   ")
        r = self.rns[ik]
        old = f"{self.data[r][col]}"
        if snapshot:
            self._ensure_snapshot("ctrl x, v, del key", self.snapshot_ctrl_x_v_del_key)
            if self.vs:
                self.vs[-1].setdefault("cells", {})[(r, col)] = old
        self.changelog_append(
            "Edit cell",
            f"ID: {self.data[r][self.ic]} column #{col + 1} named: {self.headers[col].name} with type: {self.headers[col].type_}",
            old,
            value,
        )
        self.data[r][col] = value
        return ok({"id": self.nodes[ik].name, "column": self.headers[col].name, "old": old, "new": value})

    def _get_lvls(self, iid: str, lvl=1, h=None):
        # Relative descendant buckets for delete (children of iid at `lvl`).
        # Query/CLI uses _level_index for absolute Treeview depth instead.
        if h is None:
            h = self.pc
        levels = defaultdict(list)
        stack = [(iid, lvl - 1)]
        while stack:
            current_iid, current_lvl = stack.pop()
            children = self.nodes[current_iid].cn[h]
            next_lvl = current_lvl + 1
            if next_lvl not in levels:
                levels[next_lvl] = []
            for child in children:
                levels[next_lvl].append(child)
                stack.append((child, next_lvl))
        return levels

    def _node_level(self, iid, h=None):
        h = self.pc if h is None else h
        if iid not in self.nodes or self.nodes[iid].ps[h] is None:
            return None
        level = 1
        current = iid
        while self.nodes[current].ps[h]:
            current = self.nodes[current].ps[h]
            level += 1
        return level

    def _in_subtree(self, iid, ancestor, h):
        current = iid
        while True:
            if current == ancestor:
                return True
            parent = self.nodes[current].ps[h]
            if not parent:
                return False
            current = parent

    def _level_index(self, h=None, under=None, stop_level=None):
        # Absolute 1-based depth from tops (Treeview levels / get_node_level).
        # Tree order, same walk as write_treeview. stop_level does not descend further.
        h = self.pc if h is None else int(h)
        by_level = defaultdict(list)
        levels = {}
        if under is None:
            start = [(iid, 1) for iid in self._top_iids_h(h)]
        else:
            if under not in self.nodes or self.nodes[under].ps[h] is None:
                return by_level, levels
            start = [(under, self._node_level(under, h))]
        stack = list(reversed(start))
        while stack:
            iid, lvl = stack.pop()
            if stop_level is not None and lvl > stop_level:
                continue
            by_level[lvl].append(iid)
            levels[iid] = lvl
            if stop_level is not None and lvl >= stop_level:
                continue
            for child in reversed(self.nodes[iid].cn[h]):
                stack.append((child, lvl + 1))
        return by_level, levels

    def _search_cols(self, in_, column_idx=None):
        hiers = set(self.hiers)
        if column_idx is not None:
            if column_idx == self.ic:
                role = "id"
            elif column_idx in hiers:
                role = "parent"
            else:
                role = "detail"
            if in_ == "id" and role != "id":
                return []
            if in_ == "detail" and role != "detail":
                return []
            return [column_idx]
        if in_ == "id":
            return [self.ic]
        n = len(self.headers)
        if in_ == "detail":
            return [i for i in range(n) if i != self.ic and i not in hiers]
        return list(range(n))

    @staticmethod
    def _cell_matches(cell, term_l, exact):
        cell_l = cell.lower()
        return cell_l == term_l if exact else term_l in cell_l

    def _row_matches_term(self, ik, term_l, in_, exact, column_idx=None):
        row = self.data[self.rns[ik]]
        return any(self._cell_matches(row[c], term_l, exact) for c in self._search_cols(in_, column_idx))

    def _row_find_hits(self, ik, row, term_l, in_, exact, column_idx, level):
        name = self.nodes[ik].name if ik in self.nodes else row[self.ic]
        hiers = set(self.hiers)
        hits = []
        for c in self._search_cols(in_, column_idx):
            cell = row[c]
            if not self._cell_matches(cell, term_l, exact):
                continue
            if c == self.ic:
                role = "id"
            elif c in hiers:
                role = "parent"
            else:
                role = "detail"
            hits.append(
                {
                    "id": name,
                    "column": self.headers[c].name,
                    "column_index": c,
                    "type": role,
                    "text": cell,
                    "exact": cell.lower() == term_l,
                    "level": level,
                }
            )
        return hits

    def _del_id_selection_roots(self, iids):
        selected = []
        seen = set()
        for iid in iids:
            ik = iid.lower()
            if ik in self.nodes and ik not in seen:
                selected.append(ik)
                seen.add(ik)
        selected_set = set(selected)
        roots = []
        for ik in selected:
            p = self.nodes[ik].ps[self.pc]
            while p:
                if p in selected_set:
                    break
                p = self.nodes[p].ps[self.pc] if p in self.nodes else ""
            else:
                roots.append(ik)
        return roots

    def _del_id_core(self, name: str, to_del: list[str] | None = None, snapshot: bool = True) -> list[str]:
        if to_del is None:
            to_del = []
        iid = name.lower()
        if iid not in self.nodes or self.nodes[iid].ps[self.pc] is None:
            return to_del
        pk = self.get_ids_parent(iid)
        if pk:
            self.nodes[pk].cn[self.pc].remove(iid)
        if not self.auto_sort_nodes_bool:
            if pk == "":
                self.topnodes_order[self.pc].remove(iid)
                for ciid in self.nodes[iid].cn[self.pc]:
                    self.topnodes_order[self.pc].append(ciid)
            else:
                for ciid in self.nodes[iid].cn[self.pc]:
                    self.nodes[pk].cn[self.pc].append(ciid)
        else:
            if pk:
                for ciid in self.nodes[iid].cn[self.pc]:
                    self.nodes[pk].cn[self.pc].append(ciid)
                self.nodes[pk].cn[self.pc] = self.sort_node_cn(self.nodes[pk].cn[self.pc], self.pc)
        if pk:
            for ciid in self.nodes[iid].cn[self.pc]:
                rn = self.rns[ciid]
                if snapshot and rn not in self.vs[-1]["rows"]:
                    self.vs[-1]["rows"][rn] = RowStorage(
                        0,
                        zlib.compress(pickle.dumps([self.data[rn][h] for h in self.hiers])),
                    )
                self.nodes[ciid].ps[self.pc] = pk
                self.data[rn][self.pc] = self.nodes[pk].name
                self.refresh_rows.add(ciid)
        elif pk == "":
            for ciid in self.nodes[iid].cn[self.pc]:
                rn = self.rns[ciid]
                if snapshot and rn not in self.vs[-1]["rows"]:
                    self.vs[-1]["rows"][rn] = RowStorage(
                        0,
                        zlib.compress(pickle.dumps([self.data[rn][h] for h in self.hiers])),
                    )
                self.nodes[ciid].ps[self.pc] = ""
                self.data[rn][self.pc] = ""
                self.refresh_rows.add(ciid)
        rn = self.rns[iid]
        if sum(1 for v in self.nodes[iid].ps.values() if v is not None) < 2:
            if snapshot:
                self.vs[-1]["rows"][rn] = RowStorage(1, self.data[rn])
            del self.nodes[iid]
            self._untag_id(iid)
            to_del.append(iid)
            self.refresh_rows.discard(iid)
        else:
            if snapshot and rn not in self.vs[-1]["rows"]:
                self.vs[-1]["rows"][rn] = RowStorage(
                    0,
                    zlib.compress(pickle.dumps([self.data[rn][h] for h in self.hiers])),
                )
            self.nodes[iid].cn[self.pc] = []
            self.nodes[iid].ps[self.pc] = None
            self.data[rn][self.pc] = ""
            self.refresh_rows.add(iid)
        if self.auto_sort_nodes_bool and pk and self.nodes[pk].ps[self.pc]:
            parent_parent_node = self.nodes[self.nodes[pk].ps[self.pc]]
            parent_parent_node.cn[self.pc] = self.sort_node_cn(parent_parent_node.cn[self.pc], self.pc)
        return to_del

    def _del_id_all_core(self, name: str, to_del: list[str] | None = None, snapshot: bool = True) -> list[str]:
        if to_del is None:
            to_del = []
        iid = name.lower()
        if iid not in self.nodes:
            return to_del
        self._untag_id(iid)
        to_sort = set()
        if not self.auto_sort_nodes_bool:
            for h, pk in self.nodes[iid].ps.items():
                if pk == "":
                    self.topnodes_order[h].remove(iid)
                    for ciid in self.nodes[iid].cn[h]:
                        child = self.nodes[ciid]
                        self.topnodes_order[h].append(ciid)
                        child.ps[h] = ""
                        rn = self.rns[ciid]
                        if snapshot and rn not in self.vs[-1]["rows"]:
                            self.vs[-1]["rows"][rn] = RowStorage(
                                0,
                                zlib.compress(pickle.dumps([self.data[rn][h_] for h_ in self.hiers])),
                            )
                            self.refresh_rows.add(ciid)
                        self.data[rn][h] = ""
                elif pk:
                    self.nodes[pk].cn[h].remove(iid)
                    for ciid in self.nodes[iid].cn[h]:
                        self.nodes[pk].cn[h].append(ciid)
                        child = self.nodes[ciid]
                        child.ps[h] = pk
                        rn = self.rns[ciid]
                        if snapshot and rn not in self.vs[-1]["rows"]:
                            self.vs[-1]["rows"][rn] = RowStorage(
                                0,
                                zlib.compress(pickle.dumps([self.data[rn][h_] for h_ in self.hiers])),
                            )
                            self.refresh_rows.add(ciid)
                        self.data[rn][h] = self.nodes[pk].name
        else:
            for h, pk in self.nodes[iid].ps.items():
                if pk == "":
                    for ciid in self.nodes[iid].cn[h]:
                        child = self.nodes[ciid]
                        child.ps[h] = ""
                        rn = self.rns[ciid]
                        if snapshot and rn not in self.vs[-1]["rows"]:
                            self.vs[-1]["rows"][rn] = RowStorage(
                                0,
                                zlib.compress(pickle.dumps([self.data[rn][h_] for h_ in self.hiers])),
                            )
                            self.refresh_rows.add(ciid)
                        self.data[rn][h] = ""
                elif pk:
                    self.nodes[pk].cn[h].remove(iid)
                    for ciid in self.nodes[iid].cn[h]:
                        self.nodes[pk].cn[h].append(ciid)
                        child = self.nodes[ciid]
                        child.ps[h] = pk
                        rn = self.rns[ciid]
                        if snapshot and rn not in self.vs[-1]["rows"]:
                            self.vs[-1]["rows"][rn] = RowStorage(
                                0,
                                zlib.compress(pickle.dumps([self.data[rn][h_] for h_ in self.hiers])),
                            )
                            self.refresh_rows.add(ciid)
                        self.data[rn][h] = self.nodes[pk].name
                    to_sort.add((pk, h))
                    if self.nodes[pk].ps[h]:
                        to_sort.add((self.nodes[pk].ps[h], h))
        rn = self.rns[iid]
        if snapshot:
            self.vs[-1]["rows"][rn] = RowStorage(1, self.data[rn])
        del self.nodes[iid]
        to_del.append(iid)
        self.refresh_rows.discard(iid)
        if self.auto_sort_nodes_bool:
            for node_id, h in to_sort:
                if node_id in self.nodes:
                    self.nodes[node_id].cn[h] = self.sort_node_cn(self.nodes[node_id].cn[h], h)
        return to_del

    def _del_id_orphan_core(
        self, name: str, parent: str, to_del: list[str] | None = None, snapshot: bool = True
    ) -> list[str]:
        if to_del is None:
            to_del = []
        ik = name.lower()
        if ik not in self.nodes or self.nodes[ik].ps[self.pc] is None:
            return to_del
        pk = parent.lower()
        if pk:
            self.nodes[pk].cn[self.pc].remove(ik)
        if not self.auto_sort_nodes_bool:
            if pk == "":
                self.topnodes_order[self.pc].remove(ik)
            for ciid in self.nodes[ik].cn[self.pc]:
                self.topnodes_order[self.pc].append(ciid)
        for ciid in self.nodes[ik].cn[self.pc]:
            rn = self.rns[ciid]
            child = self.nodes[ciid]
            if snapshot and rn not in self.vs[-1]["rows"]:
                self.vs[-1]["rows"][rn] = RowStorage(
                    0,
                    zlib.compress(pickle.dumps([self.data[rn][h] for h in self.hiers])),
                )
                self.refresh_rows.add(ciid)
            child.ps[self.pc] = ""
            self.data[rn][self.pc] = ""
        rn = self.rns[ik]
        if sum(1 for v in self.nodes[ik].ps.values() if v is not None) < 2:
            if snapshot:
                self.vs[-1]["rows"][rn] = RowStorage(1, self.data[rn])
            del self.nodes[ik]
            self._untag_id(ik)
            to_del.append(ik)
            self.refresh_rows.discard(ik)
        else:
            if snapshot and rn not in self.vs[-1]["rows"]:
                self.vs[-1]["rows"][rn] = RowStorage(
                    0,
                    zlib.compress(pickle.dumps([self.data[rn][h] for h in self.hiers])),
                )
                self.refresh_rows.add(ik)
            self.nodes[ik].cn[self.pc] = []
            self.nodes[ik].ps[self.pc] = None
            self.data[rn][self.pc] = ""
        if self.auto_sort_nodes_bool and pk and self.nodes[pk].ps[self.pc]:
            parent_parent_node = self.nodes[self.nodes[pk].ps[self.pc]]
            parent_parent_node.cn[self.pc] = self.sort_node_cn(parent_parent_node.cn[self.pc], self.pc)
        return to_del

    def _del_id_all_orphan_core(self, name: str, to_del: list[str] | None = None, snapshot: bool = True) -> list[str]:
        if to_del is None:
            to_del = []
        ik = name.lower()
        if ik not in self.nodes:
            return to_del
        to_sort = set()
        self._untag_id(ik)
        if not self.auto_sort_nodes_bool:
            for h, p in self.nodes[ik].ps.items():
                if p == "":
                    self.topnodes_order[h].remove(ik)
                for ciid in self.nodes[ik].cn[h]:
                    self.topnodes_order[h].append(ciid)
        for h, p in self.nodes[ik].ps.items():
            if p:
                self.nodes[p].cn[h].remove(ik)
                if self.auto_sort_nodes_bool and self.nodes[p].ps[h]:
                    to_sort.add((self.nodes[p].ps[h], h))
        for h, cn in self.nodes[ik].cn.items():
            for ciid in cn:
                rn = self.rns[ciid]
                if snapshot and rn not in self.vs[-1]["rows"]:
                    self.vs[-1]["rows"][rn] = RowStorage(
                        0,
                        zlib.compress(pickle.dumps([self.data[rn][hx] for hx in self.hiers])),
                    )
                    self.refresh_rows.add(ciid)
                child = self.nodes[ciid]
                child.ps[h] = ""
                self.data[rn][h] = ""
        rn = self.rns[ik]
        if snapshot:
            self.vs[-1]["rows"][rn] = RowStorage(1, self.data[rn])
        del self.nodes[ik]
        to_del.append(ik)
        self.refresh_rows.discard(ik)
        if self.auto_sort_nodes_bool:
            for node_id, h in to_sort:
                if node_id in self.nodes:
                    self.nodes[node_id].cn[h] = self.sort_node_cn(self.nodes[node_id].cn[h], h)
        return to_del

    def _del_id_children_core(self, name: str, to_del: list[str] | None = None, snapshot: bool = True) -> list[str]:
        if to_del is None:
            to_del = []
        ik = name.lower()
        if ik not in self.nodes or self.nodes[ik].ps[self.pc] is None:
            return to_del
        levels = self._get_lvls(ik)
        for lvl in sorted(((k, v) for k, v in levels.items()), key=itemgetter(0), reverse=True):
            for ik_ in lvl[1]:
                if ik_ not in self.nodes:
                    continue
                rn = self.rns[ik_]
                if sum(1 for v in self.nodes[ik_].ps.values() if v is not None) < 2:
                    if snapshot:
                        self.vs[-1]["rows"][rn] = RowStorage(1, self.data[rn])
                    del self.nodes[ik_]
                    to_del.append(ik_)
                    self.refresh_rows.discard(ik_)
                    self._untag_id(ik_)
                else:
                    if snapshot and rn not in self.vs[-1]["rows"]:
                        self.vs[-1]["rows"][rn] = RowStorage(
                            0,
                            zlib.compress(pickle.dumps([self.data[rn][h] for h in self.hiers])),
                        )
                        self.refresh_rows.add(ik_)
                    self.nodes[ik_].cn[self.pc] = []
                    self.nodes[ik_].ps[self.pc] = None
                    self.data[rn][self.pc] = ""
        pk = self.get_ids_parent(ik)
        rn = self.rns[ik]
        if pk:
            self.nodes[pk].cn[self.pc].remove(ik)
        if sum(1 for v in self.nodes[ik].ps.values() if v is not None) < 2:
            if snapshot:
                self.vs[-1]["rows"][rn] = RowStorage(1, self.data[rn])
            del self.nodes[ik]
            to_del.append(ik)
            self.refresh_rows.discard(ik)
            self._untag_id(ik)
        else:
            if snapshot and rn not in self.vs[-1]["rows"]:
                self.vs[-1]["rows"][rn] = RowStorage(
                    0,
                    zlib.compress(pickle.dumps([self.data[rn][h] for h in self.hiers])),
                )
                self.refresh_rows.add(ik)
            self.nodes[ik].cn[self.pc] = []
            self.nodes[ik].ps[self.pc] = None
            self.data[rn][self.pc] = ""
        if self.auto_sort_nodes_bool:
            if pk and pk in self.nodes and self.nodes[pk].ps[self.pc]:
                parent_parent_node = self.nodes[self.nodes[pk].ps[self.pc]]
                parent_parent_node.cn[self.pc] = self.sort_node_cn(parent_parent_node.cn[self.pc], self.pc)
        elif not self.auto_sort_nodes_bool and pk == "":
            try_remove(self.topnodes_order[self.pc], ik)
        return to_del

    def _del_id_children_all_core(self, name: str, to_del: list[str] | None = None, snapshot: bool = True) -> list[str]:
        if to_del is None:
            to_del = []
        ik = name.lower()
        if ik not in self.nodes or self.nodes[ik].ps[self.pc] is None:
            return to_del
        levels = self._get_lvls(ik)
        del_set = {ik, *(descendant for lvl in levels.values() for descendant in lvl)}
        to_sort = set()
        for lvl in sorted(((k, v) for k, v in levels.items()), key=itemgetter(0), reverse=True):
            for descendant in lvl[1]:
                if descendant not in self.nodes:
                    continue
                for h, p in self.nodes[descendant].ps.items():
                    if h == self.pc or p is None:
                        continue
                    if p == "":
                        if not self.auto_sort_nodes_bool:
                            self.topnodes_order[h].remove(descendant)
                    elif p not in del_set and p in self.nodes:
                        self.nodes[p].cn[h].remove(descendant)
                        if self.auto_sort_nodes_bool:
                            to_sort.add((p, h))
                    if self.auto_sort_nodes_bool and p and p in self.nodes and self.nodes[p].ps[h]:
                        gp = self.nodes[p].ps[h]
                        if gp and gp not in del_set:
                            to_sort.add((gp, h))
                    for ciid in self.nodes[descendant].cn[h]:
                        if ciid not in self.nodes or ciid in del_set:
                            continue
                        rn = self.rns[ciid]
                        if snapshot and rn not in self.vs[-1]["rows"]:
                            self.vs[-1]["rows"][rn] = RowStorage(
                                0,
                                zlib.compress(pickle.dumps([self.data[rn][h] for h in self.hiers])),
                            )
                            self.refresh_rows.add(ciid)
                        if p == "" or p in del_set or p not in self.nodes:
                            self.nodes[ciid].ps[h] = ""
                            self.data[rn][h] = ""
                            if not self.auto_sort_nodes_bool:
                                self.topnodes_order[h].append(ciid)
                        else:
                            self.nodes[ciid].ps[h] = p
                            self.nodes[p].cn[h].append(ciid)
                            self.data[rn][h] = self.nodes[p].name
                rn = self.rns[descendant]
                if snapshot:
                    self.vs[-1]["rows"][rn] = RowStorage(1, self.data[rn])
                to_del.append(descendant)
                self.refresh_rows.discard(descendant)
                self._untag_id(descendant)
                del self.nodes[descendant]

        pk = self.get_ids_parent(ik)
        rn = self.rns[ik]
        if snapshot:
            self.vs[-1]["rows"][rn] = RowStorage(1, self.data[rn])
        to_del.append(ik)
        self.refresh_rows.discard(ik)
        self._untag_id(ik)
        for h, p in self.nodes[ik].ps.items():
            if p is None:
                continue
            if p == "":
                if not self.auto_sort_nodes_bool:
                    self.topnodes_order[h].remove(ik)
            elif p not in del_set and p in self.nodes:
                self.nodes[p].cn[h].remove(ik)
                if self.auto_sort_nodes_bool:
                    to_sort.add((p, h))
            if self.auto_sort_nodes_bool and p and p in self.nodes and self.nodes[p].ps[h]:
                gp = self.nodes[p].ps[h]
                if gp and gp not in del_set:
                    to_sort.add((gp, h))
            if h == self.pc:
                continue
            for ciid in self.nodes[ik].cn[h]:
                if ciid not in self.nodes or ciid in del_set:
                    continue
                rn = self.rns[ciid]
                if snapshot and rn not in self.vs[-1]["rows"]:
                    self.vs[-1]["rows"][rn] = RowStorage(
                        0,
                        zlib.compress(pickle.dumps([self.data[rn][h] for h in self.hiers])),
                    )
                    self.refresh_rows.add(ciid)
                if p == "" or p in del_set or p not in self.nodes:
                    self.nodes[ciid].ps[h] = ""
                    self.data[rn][h] = ""
                    if not self.auto_sort_nodes_bool:
                        self.topnodes_order[h].append(ciid)
                else:
                    self.nodes[ciid].ps[h] = p
                    self.nodes[p].cn[h].append(ciid)
                    self.data[rn][h] = self.nodes[p].name
        del self.nodes[ik]
        if self.auto_sort_nodes_bool:
            for iid, h in to_sort:
                if iid in self.nodes:
                    self.nodes[iid].cn[h] = self.sort_node_cn(self.nodes[iid].cn[h], h)
            if pk and pk in self.nodes and self.nodes[pk].ps[self.pc]:
                parent_parent_node = self.nodes[self.nodes[pk].ps[self.pc]]
                parent_parent_node.cn[self.pc] = self.sort_node_cn(parent_parent_node.cn[self.pc], self.pc)
        elif not self.auto_sort_nodes_bool and pk == "":
            try_remove(self.topnodes_order[self.pc], ik)
        return to_del

    def delete(self, ids, *, children=False, all_hierarchies=False, orphan=False, snapshot=True):
        if children and orphan:
            return fail("usage", "Cannot combine children and orphan   ")
        requested = list(ids)
        result = {
            "requested": requested,
            "removed_entirely": [],
            "removed_from_hierarchy": [],
            "promoted": [],
            "deleted_descendants": [],
            "orphaned": [],
            "not_in_hierarchy": [],
            "not_found": [],
        }
        log = ChangeBuilder() if snapshot and not orphan else None
        names = []
        seen = set()
        for name in requested:
            ik = name.lower()
            if ik not in self.nodes:
                result["not_found"].append(name)
                continue
            if not all_hierarchies and self.nodes[ik].ps[self.pc] is None:
                result["not_in_hierarchy"].append(self.nodes[ik].name)
                continue
            if ik not in seen:
                names.append(self.nodes[ik].name)
                seen.add(ik)
        if children:
            names = [self.nodes[ik].name for ik in self._del_id_selection_roots(names)]
        orphan_log_id = names[0] if names else ""
        orphan_log_par = ""
        if names and names[0].lower() in self.nodes and self.nodes[names[0].lower()].ps[self.pc]:
            orphan_log_par = self.nodes[self.nodes[names[0].lower()].ps[self.pc]].name
        if snapshot and names:
            self._ensure_snapshot("delete ids", self.snapshot_delete_ids)
        self.refresh_rows = set()
        to_del = []
        processed = 0
        for name in names:
            ik = name.lower()
            if ik not in self.nodes:
                continue
            children_before = [self.nodes[c].name for c in self.nodes[ik].cn[self.pc]] if ik in self.nodes else []
            par = ""
            if self.nodes[ik].ps[self.pc]:
                par = self.nodes[self.nodes[ik].ps[self.pc]].name
            if orphan and all_hierarchies:
                to_del = self._del_id_all_orphan_core(ik, to_del, snapshot=snapshot)
                result["removed_entirely"].append(name)
                result["orphaned"].extend(children_before)
                processed += 1
            elif orphan:
                if self.nodes[ik].ps[self.pc] is None:
                    result["not_in_hierarchy"].append(name)
                    continue
                to_del = self._del_id_orphan_core(ik, par, to_del, snapshot=snapshot)
                if ik not in self.nodes:
                    result["removed_entirely"].append(name)
                else:
                    result["removed_from_hierarchy"].append(name)
                result["orphaned"].extend(children_before)
                processed += 1
            elif children and all_hierarchies:
                desc = [self.nodes[c].name for c in self.nodes[ik].cn[self.pc]]
                to_del = self._del_id_children_all_core(ik, to_del, snapshot=snapshot)
                result["deleted_descendants"].extend(desc)
                processed += 1
            elif children:
                desc = [self.nodes[c].name for c in self.nodes[ik].cn[self.pc]]
                to_del = self._del_id_children_core(ik, to_del, snapshot=snapshot)
                result["deleted_descendants"].extend(desc)
                processed += 1
            elif all_hierarchies:
                to_del = self._del_id_all_core(ik, to_del, snapshot=snapshot)
                processed += 1
            else:
                to_del = self._del_id_core(ik, to_del, snapshot=snapshot)
                result["promoted"].extend(children_before)
                processed += 1
            if log is not None:
                par_txt = par if par else "n/a - Top ID"
                col_txt = f"column #{self.pc + 1} named: {self.headers[self.pc].name}"
                if all_hierarchies and not children:
                    log.add("Delete ID from all hierarchies", f"{name}", "", "")
                elif children and all_hierarchies:
                    log.add(
                        "Delete ID + all children from all hierarchies",
                        f"ID: {name} parent: {par_txt} {col_txt}",
                        "",
                        "",
                    )
                elif children:
                    log.add("Delete ID + all children", f"ID: {name} parent: {par_txt} {col_txt}", "", "")
                else:
                    log.add("Delete ID", f"ID: {name} parent: {par_txt} {col_txt}", "", "")
        if snapshot and orphan:
            if names:
                self._log(
                    "Delete ID from all hierarchies, orphan children"
                    if all_hierarchies
                    else "Delete ID, orphan children",
                    orphan_log_id
                    if all_hierarchies
                    else (
                        f"ID: {orphan_log_id} parent: {orphan_log_par if orphan_log_par else 'n/a - Top ID'} column #{self.pc + 1} named: {self.headers[self.pc].name}"
                    ),
                    "",
                    "",
                )
        elif log is not None and processed:
            if all_hierarchies and not children:
                if len(to_del) > 1:
                    log.summary(f"Deleted {len(to_del)} IDs from all hierarchies")
                else:
                    log.force_singular = True
                    log.label = "Delete ID from all hierarchies"
            elif children and all_hierarchies:
                if processed > 1:
                    log.summary(f"Delete {processed} IDs + all children from all hierarchies")
                else:
                    log.force_singular = True
                    log.label = "Delete ID + all children from all hierarchies"
            elif children:
                if processed > 1:
                    log.summary(f"Delete {processed} IDs + all children")
                else:
                    log.force_singular = True
                    log.label = "Delete ID + all children"
            else:
                if processed > 1:
                    log.summary(f"Delete {processed} IDs")
                else:
                    log.force_singular = True
                    log.label = "Delete ID"
            if log.rows:
                self.commit_change(log)
        if to_del:
            self.ops.delete_rows([self.rns[iid] for iid in to_del if iid in self.rns])
        self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
        if not orphan:
            for name in names:
                ik = name.lower()
                if ik in to_del or ik not in self.nodes:
                    if name not in result["removed_entirely"]:
                        result["removed_entirely"].append(name)
                elif self.nodes[ik].ps[self.pc] is None and name not in result["removed_from_hierarchy"]:
                    result["removed_from_hierarchy"].append(name)
        return ok(result)

    def _push_n(self, num, sorted_seq):
        if num < sorted_seq[0]:
            return num
        hi = len(sorted_seq)
        lo = 0
        while lo < hi:
            mid = (lo + hi) // 2
            if sorted_seq[mid] < num + mid + 1:
                lo = mid + 1
            else:
                hi = mid
        return num + lo

    def _sort_later_init(self):
        if self.sort_later_dct is None:
            self.sort_later_dct = {
                "filled": False,
                "old_parents_of_parents": set(),
                "old_hier": None,
                "new_parent": (),
                "new_parent_of_parent": (),
            }

    def snapshot_paste_id(self):
        self.vs.append(
            {
                "type": "paste id",
                "rows": [],
                "required_data": self.snapshot_required_data(),
            }
        )

    def snapshot_add_col(self, treecolsel):
        self.vs.append(
            {
                "type": "add col",
                "treecolsel": int(treecolsel),
                "required_data": self.snapshot_required_data(),
            }
        )

    def snapshot_del_cols(self):
        self.vs.append(
            {
                "type": "del cols",
                "cols": {},
                "required_data": self.snapshot_required_data(),
            }
        )

    def snapshot_rename_col(self):
        self.vs.append(
            {
                "type": "rename col",
                "required_data": self.snapshot_required_data(),
            }
        )

    def _snap_parent_cells(self, rn, hier, snapshot):
        if snapshot:
            self.vs[-1]["rows"].append(
                zlib.compress(
                    pickle.dumps(
                        (
                            rn,
                            hier,
                            self.data[rn][hier],
                            self.pc,
                            self.data[rn][self.pc],
                        )
                    )
                )
            )
            self.refresh_rows.add(int(rn))

    def cut_paste(self, ID, oldparent, hier, newparent, snapshot=True, sort_later=False):
        self.refresh_rows = set()
        self._sort_later_init()
        ik = ID.lower()
        pk = oldparent.lower()
        npk = newparent.lower()
        parent_of_ik = self.nodes[ik].ps[hier]
        if ik == npk:
            return fail("cycle", "New parent is ID   ")
        if hier != self.pc and self.nodes[ik].ps[self.pc] is not None:
            return fail("already_in_hierarchy", f"ID: {ID} already in hierarchy   ")
        if npk == "":
            if self.nodes[ik].ps[self.pc] == "":
                return fail("already_in_hierarchy", f"ID: {ID} already has this parent   ")
        else:
            if self.nodes[ik].ps[self.pc] and npk == self.nodes[ik].ps[self.pc]:
                return fail("already_in_hierarchy", f"ID: {ID} already has this parent   ")
        if snapshot:
            self._ensure_snapshot("paste id", self.snapshot_paste_id)
        auto_sort_quick = self.auto_sort_nodes_bool
        for ciid in self.nodes[ik].cn[hier]:
            child = self.nodes[ciid]
            child.ps[hier] = parent_of_ik
            crow = self.rns[ciid]
            self._snap_parent_cells(crow, hier, snapshot)
            self.data[crow][hier] = self.nodes[parent_of_ik].name if parent_of_ik else ""
            if parent_of_ik:
                self.nodes[parent_of_ik].cn[hier].append(ciid)
            elif not parent_of_ik and not auto_sort_quick:
                self.topnodes_order[hier].append(ciid)
        self.nodes[ik].cn[hier] = []
        self.nodes[ik].ps[hier] = None
        if pk != "":
            self.nodes[pk].cn[hier].remove(ik)
        if npk == "":
            self.nodes[ik].ps[self.pc] = ""
        else:
            self.nodes[ik].ps[self.pc] = npk
            self.nodes[npk].cn[self.pc].append(ik)
            if auto_sort_quick:
                if sort_later and not self.sort_later_dct["filled"]:
                    self.sort_later_dct["new_parent"] = (npk, self.pc)
                    if self.nodes[npk].ps[self.pc]:
                        self.sort_later_dct["new_parent_of_parent"] = (
                            self.nodes[npk].ps[self.pc],
                            self.pc,
                        )
                elif not sort_later:
                    self.nodes[npk].cn[self.pc] = self.sort_node_cn(self.nodes[npk].cn[self.pc], self.pc)
                    if self.nodes[npk].ps[self.pc]:
                        parent_parent_node = self.nodes[self.nodes[npk].ps[self.pc]]
                        parent_parent_node.cn[self.pc] = self.sort_node_cn(parent_parent_node.cn[self.pc], self.pc)
        if not auto_sort_quick:
            if pk == "":
                try_remove(self.topnodes_order[hier], ik)
            if npk == "":
                self.topnodes_order[self.pc].append(ik)
        idrow = self.rns[ik]
        self._snap_parent_cells(idrow, hier, snapshot)
        self.data[idrow][hier] = ""
        self.data[idrow][self.pc] = newparent
        if auto_sort_quick and parent_of_ik and self.nodes[parent_of_ik].ps[hier]:
            parent_parent_iid = self.nodes[parent_of_ik].ps[hier]
            if sort_later:
                self.sort_later_dct["old_parents_of_parents"].add(parent_parent_iid)
                if self.sort_later_dct["old_hier"] is None:
                    self.sort_later_dct["old_hier"] = hier
            elif not sort_later:
                parent_parent_node = self.nodes[parent_parent_iid]
                parent_parent_node.cn[hier] = self.sort_node_cn(parent_parent_node.cn[hier], hier)
        self.sort_later_dct["filled"] = True
        return ok({})

    def cut_paste_all(self, ID, oldparent, hier, newparent, snapshot=True, sort_later=False):
        self.refresh_rows = set()
        self._sort_later_init()
        ik = ID.lower()
        pk = oldparent.lower()
        npk = newparent.lower()
        if hier != self.pc:
            if self.nodes[ik].ps[self.pc] is not None:
                return fail("already_in_hierarchy", f"ID: {ID} already in hierarchy   ")
            for ck in self.check_cn(ik, hier):
                if self.nodes[ck].ps[self.pc] is not None:
                    return fail(
                        "already_in_hierarchy",
                        f"ID: {self.nodes[ck].name} is already in hierarchy   ",
                    )
        else:
            if any(npk == ck for ck in self.check_cn(ik, hier)):
                return fail("cycle", f"Cannot add ID: {ID} to same line   ")
        if npk == "":
            if self.nodes[ik].ps[self.pc] == "":
                return fail("already_in_hierarchy", f"ID: {ID} already has this parent   ")
        else:
            if self.nodes[ik].ps[self.pc] and npk == self.nodes[ik].ps[self.pc]:
                return fail("already_in_hierarchy", f"ID: {ID} already has this parent   ")
        if snapshot:
            self._ensure_snapshot("paste id", self.snapshot_paste_id)
        self.nodes[ik].ps[hier] = None
        if pk != "":
            self.nodes[pk].cn[hier].remove(ik)
        if npk == "":
            self.nodes[ik].ps[self.pc] = ""
        else:
            self.nodes[ik].ps[self.pc] = npk
            self.nodes[npk].cn[self.pc].append(ik)
        if not self.auto_sort_nodes_bool:
            if pk == "":
                try_remove(self.topnodes_order[hier], ik)
            if npk == "":
                self.topnodes_order[self.pc].append(ik)
        idrow = self.rns[ik]
        self._snap_parent_cells(idrow, hier, snapshot)
        self.data[idrow][hier] = ""
        self.data[idrow][self.pc] = newparent
        if hier != self.pc:
            self.nodes[ik].cn[self.pc] = list(self.nodes[ik].cn[hier])
            self.nodes[ik].cn[hier] = []
            stack = [self.nodes[ik]]
            while stack:
                node = stack.pop()
                for ciid in node.cn[self.pc]:
                    child = self.nodes[ciid]
                    child.ps[self.pc] = child.ps[hier]
                    child.ps[hier] = None
                    child.cn[self.pc] = list(child.cn[hier])
                    child.cn[hier] = []
                    crow = self.rns[ciid]
                    self._snap_parent_cells(crow, hier, snapshot)
                    self.data[crow][self.pc] = f"{self.data[crow][hier]}"
                    self.data[crow][hier] = ""
                    stack.append(child)
        if self.auto_sort_nodes_bool:
            if sort_later:
                if npk and not self.sort_later_dct["filled"]:
                    self.sort_later_dct["new_parent"] = (npk, self.pc)
                    if self.nodes[npk].ps[self.pc]:
                        self.sort_later_dct["new_parent_of_parent"] = (
                            self.nodes[npk].ps[self.pc],
                            self.pc,
                        )
                if pk and self.nodes[pk].ps[hier]:
                    self.sort_later_dct["old_parents_of_parents"].add(self.nodes[pk].ps[hier])
                    if self.sort_later_dct["old_hier"] is None:
                        self.sort_later_dct["old_hier"] = hier
            elif not sort_later:
                if npk:
                    self.nodes[npk].cn[self.pc] = self.sort_node_cn(self.nodes[npk].cn[self.pc], self.pc)
                    if self.nodes[npk].ps[self.pc]:
                        parent_parent_node = self.nodes[self.nodes[npk].ps[self.pc]]
                        parent_parent_node.cn[self.pc] = self.sort_node_cn(parent_parent_node.cn[self.pc], self.pc)
                if pk and self.nodes[pk].ps[hier]:
                    parent_parent_node = self.nodes[self.nodes[pk].ps[hier]]
                    parent_parent_node.cn[hier] = self.sort_node_cn(parent_parent_node.cn[hier], hier)
        self.sort_later_dct["filled"] = True
        return ok({})

    def copy_paste(self, ID, hier, newparent, snapshot=True, sort_later=False):
        self.refresh_rows = set()
        self._sort_later_init()
        ik = ID.lower()
        npk = newparent.lower()
        if hier == self.pc or self.nodes[ik].ps[self.pc] is not None:
            return fail("already_in_hierarchy", f"ID {ID} already in hierarchy   ")
        if snapshot:
            self._ensure_snapshot("paste id", self.snapshot_paste_id)
        if npk == "":
            self.nodes[ik].ps[self.pc] = ""
        else:
            self.nodes[ik].ps[self.pc] = npk
            self.nodes[npk].cn[self.pc].append(ik)
            if self.auto_sort_nodes_bool:
                if sort_later and not self.sort_later_dct["filled"]:
                    self.sort_later_dct["new_parent"] = (npk, self.pc)
                    if self.nodes[npk].ps[self.pc]:
                        self.sort_later_dct["new_parent_of_parent"] = (
                            self.nodes[npk].ps[self.pc],
                            self.pc,
                        )
                elif not sort_later:
                    self.nodes[npk].cn[self.pc] = self.sort_node_cn(self.nodes[npk].cn[self.pc], self.pc)
                    if self.nodes[npk].ps[self.pc]:
                        parent_parent_node = self.nodes[self.nodes[npk].ps[self.pc]]
                        parent_parent_node.cn[self.pc] = self.sort_node_cn(parent_parent_node.cn[self.pc], self.pc)
        if not self.auto_sort_nodes_bool and npk == "":
            self.topnodes_order[self.pc].append(ik)
        rn = self.rns[ik]
        self._snap_parent_cells(rn, hier, snapshot)
        self.data[rn][self.pc] = newparent
        self.sort_later_dct["filled"] = True
        return ok({})

    def copy_paste_all(self, ID, hier, newparent, snapshot=True, sort_later=False):
        self.refresh_rows = set()
        self._sort_later_init()
        ik = ID.lower()
        npk = newparent.lower()
        if hier == self.pc or self.nodes[ik].ps[self.pc] is not None:
            return fail("already_in_hierarchy", f"ID {ID} already in hierarchy   ")
        for ck in self.check_cn(ik, hier):
            if self.nodes[ck].ps[self.pc] is not None:
                return fail("already_in_hierarchy", f"ID: {self.nodes[ck].name} is already in hierarchy   ")
        if snapshot:
            self._ensure_snapshot("paste id", self.snapshot_paste_id)
        if npk == "":
            self.nodes[ik].ps[self.pc] = ""
        else:
            self.nodes[ik].ps[self.pc] = npk
            self.nodes[npk].cn[self.pc].append(ik)
        if not self.auto_sort_nodes_bool and npk == "":
            self.topnodes_order[self.pc].append(ik)
        rn = self.rns[ik]
        self._snap_parent_cells(rn, hier, snapshot)
        self.data[rn][self.pc] = newparent
        self.nodes[ik].cn[self.pc] = list(self.nodes[ik].cn[hier])
        stack = list(self.nodes[ik].cn[hier])
        while stack:
            ciid = stack.pop()
            child = self.nodes[ciid]
            crow = self.rns[ciid]
            self._snap_parent_cells(crow, hier, snapshot)
            child.ps[self.pc] = child.ps[hier]
            child.cn[self.pc] = list(child.cn[hier])
            self.data[crow][self.pc] = f"{self.data[crow][hier]}"
            stack.extend(child.cn[hier])
        if npk and self.auto_sort_nodes_bool:
            if sort_later and not self.sort_later_dct["filled"]:
                self.sort_later_dct["new_parent"] = (npk, self.pc)
                if self.nodes[npk].ps[self.pc]:
                    self.sort_later_dct["new_parent_of_parent"] = (
                        self.nodes[npk].ps[self.pc],
                        self.pc,
                    )
            elif not sort_later:
                self.nodes[npk].cn[self.pc] = self.sort_node_cn(self.nodes[npk].cn[self.pc], self.pc)
                if self.nodes[npk].ps[self.pc]:
                    parent_parent_node = self.nodes[self.nodes[npk].ps[self.pc]]
                    parent_parent_node.cn[self.pc] = self.sort_node_cn(parent_parent_node.cn[self.pc], self.pc)
        self.sort_later_dct["filled"] = True
        return ok({})

    def _resolve_hier_index(self, token):
        if token is None:
            return self.pc
        if isinstance(token, int):
            if token not in self.hiers:
                return None
            return token
        t = str(token)
        for h in self.hiers:
            if self.headers[h].name.lower() == t.lower():
                return h
        return None

    def move(self, ID, *, parent=None, top=False, hier=None, from_hier=None, before=None, after=None, snapshot=True):
        ik = ID.lower()
        if ik not in self.nodes:
            return fail("id_not_found", "ID doesn't exist   ")
        dest = self._resolve_hier_index(hier)
        src = self._resolve_hier_index(from_hier)
        if dest is None:
            return fail("unknown_hierarchy", "Unknown hierarchy   ")
        if src is None:
            return fail("unknown_hierarchy", "Unknown hierarchy   ")
        if before is not None and after is not None:
            return fail("usage", "Use only one of before or after   ")
        if top and parent:
            return fail("usage", "Use only one of parent or top   ")
        if (before is not None or after is not None) and (parent or top):
            return fail("usage", "before/after cannot be combined with parent or top   ")
        saved_pc = self.pc
        self.pc = dest
        try:
            if self.nodes[ik].ps[src] is None:
                return fail("not_in_hierarchy", f"{ID} is not in this hierarchy   ")
            if before is not None or after is not None:
                if self.auto_sort_nodes_bool:
                    return fail("auto_sort_on", "Turn auto-sort off to use before/after   ")
                sib = before if before is not None else after
                sk = sib.lower()
                if sk not in self.nodes:
                    return fail("id_not_found", "ID doesn't exist   ", id=sib)
                if self.nodes[sk].ps[dest] is None:
                    return fail("not_in_hierarchy", f"{sib} is not in this hierarchy   ")
                implied_key = self.nodes[sk].ps[dest]
                parent = "" if implied_key == "" else self.nodes[implied_key].name
                top = False
            if top:
                parent = ""
            if parent is None:
                return fail("usage", "Need a parent, top, before, or after   ")
            old_pk = self.nodes[ik].ps[src]
            oldparent = "" if not old_pk else self.nodes[old_pk].name
            if snapshot:
                self._ensure_snapshot("paste id", self.snapshot_paste_id)
            out = self.cut_paste(ID, oldparent, src, parent, snapshot=snapshot)
            if not out["ok"]:
                return out
            if before is not None or after is not None:
                self._splice_tree_order(ik, before if before is not None else after, before is not None)
            if snapshot:
                display = self.nodes[ik].name
                self._log(
                    "Cut and paste ID",
                    display,
                    f"Old parent: {oldparent if oldparent else 'n/a - Top ID'} old column #{src + 1} named: {self.headers[src].name}",
                    f"New parent: {parent if parent else 'n/a - Top ID'} new column #{self.pc + 1} named: {self.headers[self.pc].name}",
                )
            return ok(
                {
                    "moved": [
                        {
                            "id": self.nodes[ik].name,
                            "from_hierarchy": self.headers[src].name,
                            "from_parent": oldparent,
                            "to_hierarchy": self.headers[dest].name,
                            "to_parent": parent,
                        }
                    ]
                }
            )
        finally:
            self.pc = saved_pc

    def copy(self, ID, *, parent=None, top=False, hier, snapshot=True):
        ik = ID.lower()
        if ik not in self.nodes:
            return fail("id_not_found", "ID doesn't exist   ")
        dest = self._resolve_hier_index(hier)
        if dest is None:
            return fail("unknown_hierarchy", "Unknown hierarchy   ")
        if top and parent:
            return fail("usage", "Use only one of parent or top   ")
        if not top and parent is None:
            return fail("usage", "Need a parent or top   ")
        if top:
            parent = ""
        src = self.pc
        saved_pc = self.pc
        self.pc = dest
        try:
            if snapshot:
                self._ensure_snapshot("paste id", self.snapshot_paste_id)
            out = self.copy_paste(ID, src, parent, snapshot=snapshot)
            if not out["ok"]:
                return out
            if snapshot:
                display = self.nodes[ik].name
                self._log(
                    "Copy and paste ID",
                    display,
                    f"From column #{src + 1} named: {self.headers[src].name}",
                    f"New parent: {parent if parent else 'n/a - Top ID'} new column #{self.pc + 1} named: {self.headers[self.pc].name}",
                )
            return ok(
                {
                    "copied": [
                        {
                            "id": self.nodes[ik].name,
                            "from_hierarchy": self.headers[src].name,
                            "to_hierarchy": self.headers[dest].name,
                            "to_parent": parent,
                        }
                    ]
                }
            )
        finally:
            self.pc = saved_pc

    def repair_empty_hierarchies(self):
        first_hier = self.hiers[0]
        quick_hiers = self.hiers[1:]
        lh = len(self.hiers)
        for node in self.nodes.values():
            tlly = 0
            for k, v in node.cn.items():
                if not (v or node.ps[k]):
                    node.ps[k] = None
                    tlly += 1
            if tlly == lh:
                node.ps[first_hier] = ""
                for h in quick_hiers:
                    node.ps[h] = None
        if not self.auto_sort_nodes_bool:
            current_nodes = dict.fromkeys(self.topnodes_order[self.hiers[0]])
            wc = []
            woc = []
            for iid, node in self.nodes.items():
                if iid not in current_nodes and node.ps[self.hiers[0]] == "":
                    if node.cn[self.hiers[0]]:
                        wc.append(iid)
                    else:
                        woc.append(iid)
            self.topnodes_order[self.hiers[0]] = (
                list(current_nodes) + sorted(wc, key=sort_key) + sorted(woc, key=sort_key)
            )

    def adjust_hiers_del_cols(self, cols):
        auto_sort_nodes_bool = self.auto_sort_nodes_bool
        self.hiers = [k if not (num := bisect_left(cols, k)) else k - num for k in self.hiers]
        for node in self.nodes.values():
            node.ps = {k if not (num := bisect_left(cols, k)) else k - num: v for k, v in node.ps.items()}
            node.cn = {k if not (num := bisect_left(cols, k)) else k - num: v for k, v in node.cn.items()}
        if not auto_sort_nodes_bool:
            self.topnodes_order = {
                k if not (num := bisect_left(cols, k)) else k - num: v for k, v in self.topnodes_order.items()
            }

    def adjust_hiers_add_cols(self, cols):
        auto_sort_nodes_bool = self.auto_sort_nodes_bool
        self.hiers = [self._push_n(k, cols) for k in self.hiers]
        for node in self.nodes.values():
            node.ps = {self._push_n(k, cols): v for k, v in node.ps.items()}
            node.cn = {self._push_n(k, cols): v for k, v in node.cn.items()}
        if not auto_sort_nodes_bool:
            self.topnodes_order = {self._push_n(k, cols): v for k, v in self.topnodes_order.items()}

    def rename_col(self, col, name, snapshot=True):
        if snapshot:
            self._ensure_snapshot("rename col", self.snapshot_rename_col)
            self.changelog_append(
                "Column rename",
                f"Column #{col + 1} with type: {self.headers[col].type_}",
                f"{self.headers[col].name}",
                f"{name}",
            )
        old = self.headers[col].name
        self.headers[col].name = name
        return ok({"old": old, "new": name, "index": col})

    def add_col(self, col, name, type_, snapshot=True):
        if snapshot:
            self._ensure_snapshot("add col", lambda: self.snapshot_add_col(col))
        self.ic = self._push_n(self.ic, [col])
        self.pc = self._push_n(self.pc, [col])
        self.row_len += 1
        self.headers.insert(col, Header(name, type_))
        self.ops.insert_cols(col, 1)
        self.adjust_hiers_add_cols(cols=[col])
        if snapshot:
            self.changelog_append(
                "Add new detail column",
                f"Column #{col} with name: {name} and type: {type_}",
                "",
                "",
            )
        return ok(
            {
                "name": name,
                "index": col,
                "letter": xl_column_string(col + 1),
                "role": "detail",
                "type": self.headers[col].type_,
            }
        )

    def add_hier_col(self, col, name, snapshot=True):
        if snapshot:
            self._ensure_snapshot("add col", lambda: self.snapshot_add_col(col))
        self.ic = self._push_n(self.ic, [col])
        self.pc = self._push_n(self.pc, [col])
        self.row_len += 1
        self.adjust_hiers_add_cols(cols=[col])
        self.hiers = sorted([col] + self.hiers)
        self.headers.insert(col, Header(name, "Parent"))
        self.ops.insert_cols(col, 1)
        for node in self.nodes.values():
            node.ps[col] = None
            node.cn[col] = []
        if not self.auto_sort_nodes_bool:
            self.topnodes_order[col] = []
        if snapshot:
            self.changelog_append(
                "Add new hierarchy column",
                f"Column #{col + 1} with name: {name}",
                "",
                "",
            )
        return ok(
            {
                "name": name,
                "index": col,
                "letter": xl_column_string(col + 1),
                "role": "parent",
                "type": "Parent",
            }
        )

    def del_cols(self, cols, snapshot=True):
        cols = list(cols)
        if snapshot:
            self._ensure_snapshot("del cols", self.snapshot_del_cols)
            cols_dict = self.vs[-1]["cols"]
            for datacn in reversed(cols):
                for rn in range(len(self.data)):
                    if datacn not in cols_dict:
                        cols_dict[datacn] = {}
                    try:
                        cols_dict[datacn][rn] = self.data[rn][datacn]
                    except Exception:
                        continue
        self.ops.delete_cols(cols)
        self.ic = self.ic if not (num := bisect_left(cols, self.ic)) else self.ic - num
        self.pc = self.pc if not (num := bisect_left(cols, self.pc)) else self.pc - num
        deleted = [{"name": self.headers[c].name, "index": c} for c in cols]
        if snapshot:
            self.changelog_append(
                "Delete columns",
                f"Columns: {', '.join(item['name'] for item in deleted)}",
                "",
                "",
            )
        cols_set = set(cols)
        self.headers = [hdr for i, hdr in enumerate(self.headers) if i not in cols_set]
        hiers_orig = self.hiers.copy()
        self.hiers = list(filterfalse(cols_set.__contains__, self.hiers))
        if hiers_to_del := list(filter(cols_set.__contains__, reversed(hiers_orig))):
            for col in hiers_to_del:
                for node in self.nodes.values():
                    del node.ps[col]
                    del node.cn[col]
                if not self.auto_sort_nodes_bool:
                    del self.topnodes_order[col]
            self.repair_empty_hierarchies()
        self.row_len -= len(cols)
        self.adjust_hiers_del_cols(cols)
        return ok({"deleted": deleted, "not_found": []})

    def top_iids(self):
        pc = self.pc
        if self.auto_sort_nodes_bool:
            wc = []
            woc = []
            for iid, node in self.nodes.items():
                if node.ps[pc] == "":
                    if node.cn[pc]:
                        wc.append(iid)
                    else:
                        woc.append(iid)
            yield from sorted(wc, key=sort_key)
            yield from sorted(woc, key=sort_key)
        else:
            yield from self.topnodes_order[pc]

    def sort_all_children(self):
        for n in self.nodes.values():
            for h, cn in n.cn.items():
                if cn:
                    n.cn[h] = self.sort_node_cn(cn, h)

    def snapshot_sheet_sort(self):
        self.vs.append(
            {
                "type": "sort",
                "ids": {v: k for k, v in self.rns.items()},
                "required_data": self.snapshot_required_data(),
            }
        )

    def snapshot_edit_validation(self, col, validation):
        self.changelog_append(
            "Edit validation",
            f"Column #{col + 1} named: {self.headers[col].name} with type: {self.headers[col].type_}",
            f"{','.join(self.headers[col].validation)}",
            f"{','.join(validation)}",
        )
        self.vs.append(
            {
                "type": "edit validation",
                "col_num": col,
                "col": zlib.compress(pickle.dumps([r[col] for r in self.data])),
                "required_data": self.snapshot_required_data(),
            }
        )

    def sort_sheet(self, header, order, snapshot=True):
        try:
            col = next(i for i, h in enumerate(self.headers) if h.name.lower() == header.lower())
        except StopIteration:
            return fail("unknown_column", f"Unknown column {header}   ")
        header = self.headers[col].name
        if snapshot:
            self._ensure_snapshot("sort", self.snapshot_sheet_sort)
            self.changelog_append(
                "Sort sheet",
                f"Sorted sheet by column #{col + 1} named: {header} in {order} order",
                "",
                "",
            )

        def ak(row):
            return tuple(int(c) if c.isdigit() else c.lower() for c in re.split("([0-9]+)", row[col]))

        if order == "ASCENDING":
            self.data.sort(key=ak)
        elif order == "DESCENDING":
            self.data.sort(key=ak, reverse=True)
        self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
        return ok({"kind": "column", "column": header, "descending": order == "DESCENDING"})

    def _walk_current_hier_rows(self):
        new_sheet = []
        visited = set()
        stack = [(iid, self.nodes[iid].cn[self.pc]) for iid in self.top_iids()]
        stack.reverse()
        while stack:
            iid, children = stack.pop()
            rowno = self.rns[iid]
            if rowno not in visited:
                visited.add(rowno)
                new_sheet.append(self.data[rowno])
                child_stack = [(ciid, self.nodes[ciid].cn[self.pc]) for ciid in reversed(children)]
                stack.extend(child_stack)
        for r in sorted(r for r in self.rns.values() if r not in visited):
            new_sheet.append(self.data[r])
        return new_sheet

    def sort_sheet_walk(self, snapshot=True):
        if snapshot:
            self._ensure_snapshot("sort", self.snapshot_sheet_sort)
            self.changelog_append(
                "Sort sheet",
                "Sorted sheet in tree walk order",
                "",
                "",
            )
        oldpc = int(self.pc)
        for h in reversed(self.hiers):
            self.pc = int(h)
            self.data = self._walk_current_hier_rows()
            self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
        self.pc = int(oldpc)
        self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
        return ok({"kind": "tree"})

    def sort_children(self, ID, snapshot=True):
        ik = ID.lower()
        if ik not in self.nodes:
            return fail("id_not_found", "ID doesn't exist   ")
        self.nodes[ik].cn[self.pc] = self.sort_node_cn(self.nodes[ik].cn[self.pc], self.pc)
        return ok({"kind": "children", "id": self.nodes[ik].name})

    def options_dict(self):
        return {
            "allow-spaces-ids": bool(self.allow_spaces_ids_var),
            "allow-spaces-columns": bool(self.allow_spaces_columns_var),
            "auto-sort": bool(self.auto_sort_nodes_bool),
            "save-program-data": bool(self.save_xlsx_with_program_data and self.save_json_with_program_data),
        }

    def set_option(self, name, value):
        key = str(name).lower()
        if isinstance(value, str):
            raw = value.lower()
            if raw not in ("on", "off"):
                return fail("unknown_option", f"Unknown option value {value}   ")
            value = raw == "on"
        elif not isinstance(value, bool):
            return fail("unknown_option", f"Unknown option value {value}   ")
        if key == "allow-spaces-ids":
            self.allow_spaces_ids_var = value
        elif key == "allow-spaces-columns":
            self.allow_spaces_columns_var = value
        elif key == "auto-sort":
            self.auto_sort_nodes_bool = value
            if value:
                self.sort_all_children()
            else:
                self.remake_topnodes_order()
        elif key == "save-program-data":
            self.save_xlsx_with_program_data = value
            self.save_json_with_program_data = value
        else:
            return fail("unknown_option", f"Unknown option {name}   ")
        return ok(self.options_dict())

    def tag(self, names, *, snapshot=False):
        tagged = []
        already = []
        not_found = []
        for name in names:
            ik = name.lower()
            if ik not in self.nodes:
                not_found.append(name)
                continue
            display = self.nodes[ik].name
            if ik in self.tagged_ids:
                already.append(display)
            else:
                self.tagged_ids.add(ik)
                tagged.append(display)
        return ok({"tagged": tagged, "already": already, "not_found": not_found})

    def untag(self, names):
        untagged = []
        not_tagged = []
        not_found = []
        for name in names:
            ik = name.lower()
            if ik not in self.nodes:
                not_found.append(name)
                continue
            display = self.nodes[ik].name
            if ik in self.tagged_ids:
                self.tagged_ids.discard(ik)
                untagged.append(display)
            else:
                not_tagged.append(display)
        return ok({"untagged": untagged, "not_tagged": not_tagged, "not_found": not_found})

    def clear_tags(self):
        n = len(self.tagged_ids)
        self.tagged_ids = set()
        return ok({"cleared": n})

    def tags_list(self):
        names = [self.nodes[ik].name for ik in self.tagged_ids if ik in self.nodes]
        names.sort(key=sort_key)
        return ok({"ids": names})

    def tag_from_terms(self, terms, *, in_="any", exact=False):
        terms_set = {str(t).lower() for t in terms if str(t).strip()}
        terms_total = len(terms_set)
        if not terms_set:
            return ok({"terms_matched": 0, "terms_total": 0, "ids_tagged": []})
        find_ids = in_ in ("id", "any")
        find_details = in_ in ("detail", "any")
        idcol_hiers = set(self.hiers) | {self.ic}
        matched_lower = set()
        ids_to_tag = []
        for r in self.data:
            row_matched = False
            if find_ids:
                cell = r[self.ic].lower()
                if exact:
                    if cell in terms_set:
                        matched_lower.add(cell)
                        row_matched = True
                else:
                    for term in terms_set:
                        if term in cell:
                            matched_lower.add(term)
                            row_matched = True
            if find_details:
                for i, e in enumerate(r):
                    if i in idcol_hiers:
                        continue
                    cell = e.lower()
                    if exact:
                        if cell in terms_set:
                            matched_lower.add(cell)
                            row_matched = True
                    else:
                        for term in terms_set:
                            if term in cell:
                                matched_lower.add(term)
                                row_matched = True
            if row_matched:
                ik = r[self.ic].lower()
                if ik in self.rns:
                    ids_to_tag.append(self.nodes[ik].name if ik in self.nodes else r[self.ic])
        tagged_out = self.tag(ids_to_tag)
        ids_tagged = tagged_out["result"]["tagged"] + tagged_out["result"]["already"]
        return ok(
            {
                "terms_matched": len(matched_lower),
                "terms_total": terms_total,
                "ids_tagged": ids_tagged,
            }
        )

    def check_validation_validity(self, col: int, validation: list[str]):
        if not validation:
            return validation
        if self.headers[col].type_ != "Text":
            return "Error: Only Detail columns can have validation"
        return validation if "" in validation else [""] + validation

    def apply_validation_to_col(self, col):
        validset = set(self.headers[col].validation)
        for rn in range(len(self.data)):
            if not self.is_in_validation(validset, self.data[rn][col]):
                self.data[rn][col] = ""

    def set_validation(self, col, values, snapshot=True):
        if not (0 <= col < len(self.headers)):
            return fail("unknown_column", "Unknown column   ")
        if values:
            validation = self.check_validation_validity(col, list(values))
            if isinstance(validation, str):
                return fail("validation", validation)
        else:
            validation = []
        if validation == self.headers[col].validation:
            return ok(
                {
                    "column": self.headers[col].name,
                    "index": col,
                    "validation": list(self.headers[col].validation),
                }
            )
        if snapshot:
            self._ensure_snapshot(
                "edit validation",
                lambda: self.snapshot_edit_validation(col, validation),
            )
        self.headers[col].validation = validation
        if validation:
            self.apply_validation_to_col(col)
        return ok(
            {
                "column": self.headers[col].name,
                "index": col,
                "validation": list(validation),
            }
        )

    def resolve_column(self, token):
        t = str(token)
        for i, header in enumerate(self.headers):
            if header.name.lower() == t.lower():
                return ok({"index": i, "name": header.name})
        if t.isdigit():
            i = int(t)
            if 0 <= i < len(self.headers):
                return ok({"index": i, "name": self.headers[i].name})
            return fail("unknown_column", f"Unknown column {token}   ")
        i = _alpha_to_idx(t)
        if i is not None and 0 <= i < len(self.headers):
            return ok({"index": i, "name": self.headers[i].name})
        return fail("unknown_column", f"Unknown column {token}   ")

    def resolve_hier(self, token):
        out = self.resolve_column(token)
        if not out["ok"]:
            return fail("unknown_hierarchy", f"Unknown hierarchy {token}   ")
        i = out["result"]["index"]
        if self.headers[i].type_ != "Parent":
            return fail("unknown_hierarchy", f"Unknown hierarchy {token}   ")
        return ok({"index": i, "name": self.headers[i].name})

    def resolve_id(self, token):
        ik = str(token).lower()
        if ik not in self.nodes:
            return fail("id_not_found", "ID doesn't exist   ")
        return ok({"key": ik, "id": self.nodes[ik].name})

    def columns_list(self):
        return ok({"columns": [self._column_entry(i) for i in range(len(self.headers))]})

    def hier_get(self):
        if not self.headers or self.pc not in self.hiers:
            return ok({"hierarchy": None, "index": None, "letter": None})
        header = self.headers[self.pc]
        return ok(
            {
                "hierarchy": header.name,
                "index": self.pc,
                "letter": xl_column_string(self.pc + 1),
            }
        )

    def hier_list(self):
        current = self.headers[self.pc].name if self.headers and self.pc in self.hiers else None
        return ok(
            {
                "hierarchies": [
                    {
                        "index": h,
                        "letter": xl_column_string(h + 1),
                        "name": self.headers[h].name,
                    }
                    for h in self.hiers
                    if h < len(self.headers)
                ],
                "current": current,
            }
        )

    def set_hier(self, token):
        out = self.resolve_hier(token)
        if not out["ok"]:
            return out
        self.pc = out["result"]["index"]
        return self.hier_get()

    def _id_object(self, ik):
        node = self.nodes[ik]
        parents = {}
        children = {}
        for h in self.hiers:
            name = self.headers[h].name
            pk = node.ps[h]
            if pk is None:
                parents[name] = None
            elif pk == "":
                parents[name] = ""
            else:
                parents[name] = self.nodes[pk].name
            children[name] = [self.nodes[c].name for c in node.cn[h]]
        details = {}
        row = self.data[self.rns[ik]]
        for i, header in enumerate(self.headers):
            if i != self.ic and i not in self.hiers:
                details[header.name] = row[i]
        return {
            "id": node.name,
            "parents": parents,
            "children": children,
            "details": details,
            "tagged": ik in self.tagged_ids,
        }

    def get_ids(self, names):
        found = []
        not_found = []
        for name in names:
            ik = name.lower()
            if ik not in self.nodes:
                not_found.append(name)
            else:
                found.append(self._id_object(ik))
        if not found:
            return fail("id_not_found", "ID doesn't exist   ")
        if len(names) == 1:
            return ok(found[0])
        return ok({"ids": found, "not_found": not_found})

    def get_query(
        self,
        names=None,
        *,
        level=None,
        contains=None,
        in_="detail",
        exact=False,
        column=None,
        under=None,
        hier=None,
        limit=50,
        ids_only=False,
    ):
        if in_ not in ("id", "detail", "any"):
            return fail("usage", "in must be id, detail, or any   ")
        if level is not None and level < 1:
            return fail("usage", "level must be an integer >= 1   ")
        if contains is not None and not str(contains):
            return fail("usage", "--contains TERM is empty   ")
        h = self.pc if hier is None else int(hier)
        if h not in self.hiers:
            return fail("unknown_hierarchy", "Unknown hierarchy   ")
        if under is not None:
            if under not in self.nodes:
                return fail("id_not_found", "ID doesn't exist   ")
            if self.nodes[under].ps[h] is None:
                return fail("not_in_hierarchy", f"{self.nodes[under].name} is not in this hierarchy   ")
        not_found = []
        if names:
            named = []
            seen = set()
            for name in names:
                ik = name.lower()
                if ik not in self.nodes:
                    not_found.append(name)
                    continue
                if ik not in seen:
                    named.append(ik)
                    seen.add(ik)
            candidates = []
            cand_levels = {}
            for ik in named:
                if self.nodes[ik].ps[h] is None:
                    continue
                if under is not None and not self._in_subtree(ik, under, h):
                    continue
                lvl = self._node_level(ik, h)
                if level is not None and lvl != level:
                    continue
                candidates.append(ik)
                cand_levels[ik] = lvl
        else:
            by_level, cand_levels = self._level_index(h, under=under, stop_level=level)
            if level is not None:
                candidates = by_level.get(level, [])
            else:
                candidates = [iid for lvl in sorted(by_level) for iid in by_level[lvl]]
        term_l = None if contains is None else str(contains).lower()
        matched = []
        for iid in candidates:
            if term_l is not None and not self._row_matches_term(iid, term_l, in_, exact, column):
                continue
            matched.append(iid)
        total = len(matched)
        shown = matched if not limit else matched[:limit]
        ids_out = []
        for iid in shown:
            lvl = cand_levels.get(iid)
            if ids_only:
                ids_out.append({"id": self.nodes[iid].name, "level": lvl})
            else:
                obj = self._id_object(iid)
                obj["level"] = lvl
                ids_out.append(obj)
        return ok(
            {
                "ids": ids_out,
                "not_found": not_found,
                "truncated": len(shown) < total,
                "returned": len(shown),
                "total": total,
                "hierarchy": self.headers[h].name,
                "level": level,
                "under": None if under is None else self.nodes[under].name,
            }
        )

    def _node_details(self, ik):
        row = self.data[self.rns[ik]]
        skip = set(self.hiers) | {self.ic}
        return {self.headers[i].name: row[i] for i in range(len(self.headers)) if i not in skip}

    def _top_iids_h(self, h):
        if self.auto_sort_nodes_bool:
            wc = []
            woc = []
            for iid, node in self.nodes.items():
                if node.ps[h] == "":
                    if node.cn[h]:
                        wc.append(iid)
                    else:
                        woc.append(iid)
            return sorted(wc, key=sort_key) + sorted(woc, key=sort_key)
        return list(self.topnodes_order.get(h, []))

    def _count_tree_nodes(self, roots, h, depth):
        n = 0
        stack = [(iid, 0) for iid in reversed(roots)]
        while stack:
            iid, d = stack.pop()
            n += 1
            if depth is None or d < depth:
                for child in reversed(self.nodes[iid].cn[h]):
                    stack.append((child, d + 1))
        return n

    def tree_dump(self, *, under=None, depth=1, details=True, hier=None, cap=500, force=False):
        h = self.pc if hier is None else int(hier)
        if h not in self.hiers:
            return fail("unknown_hierarchy", "Unknown hierarchy   ")
        if under is None:
            roots = self._top_iids_h(h)
            under_name = None
        else:
            ik = under.lower()
            if ik not in self.nodes:
                return fail("id_not_found", "ID doesn't exist   ")
            if self.nodes[ik].ps[h] is None:
                return fail("not_in_hierarchy", f"{under} is not in this hierarchy   ")
            roots = [ik]
            under_name = self.nodes[ik].name
        total = self._count_tree_nodes(roots, h, depth)
        remaining = None if force else cap
        nodes_out = []

        def walk(iid, remaining_depth):
            nonlocal remaining
            node = {"id": self.nodes[iid].name, "children": []}
            if details:
                node["details"] = self._node_details(iid)
            if remaining is not None:
                remaining -= 1
            if remaining_depth is not None and remaining_depth <= 0:
                return node
            for child in self.nodes[iid].cn[h]:
                if remaining is not None and remaining <= 0:
                    break
                child_depth = None if remaining_depth is None else remaining_depth - 1
                node["children"].append(walk(child, child_depth))
            return node

        for iid in roots:
            if remaining is not None and remaining <= 0:
                break
            nodes_out.append(walk(iid, depth))
        returned = 0
        stack = list(nodes_out)
        while stack:
            n = stack.pop()
            returned += 1
            stack.extend(n["children"])
        return ok(
            {
                "hierarchy": self.headers[h].name,
                "depth": depth,
                "under": under_name,
                "truncated": returned < total,
                "returned": returned,
                "total": total,
                "nodes": nodes_out,
            }
        )

    def find(
        self,
        term,
        *,
        in_="any",
        exact=False,
        hier=None,
        all_hier=False,
        limit=50,
        level=None,
        under=None,
        column=None,
    ):
        if in_ not in ("id", "detail", "any"):
            return fail("usage", "in must be id, detail, or any   ")
        if all_hier and (level is not None or under is not None):
            return fail("usage", "Cannot combine all-hier with level or under   ")
        if level is not None and level < 1:
            return fail("usage", "level must be an integer >= 1   ")
        term_l = str(term).lower()
        hfilter = None if all_hier else (self.pc if hier is None else int(hier))
        if hfilter is not None and hfilter not in self.hiers:
            return fail("unknown_hierarchy", "Unknown hierarchy   ")
        if under is not None:
            if under not in self.nodes:
                return fail("id_not_found", "ID doesn't exist   ")
            if hfilter is None or self.nodes[under].ps[hfilter] is None:
                return fail("not_in_hierarchy", f"{self.nodes[under].name} is not in this hierarchy   ")
        hits = []
        if hfilter is not None and (level is not None or under is not None):
            by_level, levels = self._level_index(hfilter, under=under, stop_level=level)
            if level is not None:
                iids = by_level.get(level, [])
            else:
                iids = [iid for lvl in sorted(by_level) for iid in by_level[lvl]]
            for iid in iids:
                row = self.data[self.rns[iid]]
                hits.extend(self._row_find_hits(iid, row, term_l, in_, exact, column, levels.get(iid)))
        else:
            for row in self.data:
                ik = row[self.ic].lower()
                if hfilter is not None and (ik not in self.nodes or self.nodes[ik].ps[hfilter] is None):
                    continue
                hits.extend(self._row_find_hits(ik, row, term_l, in_, exact, column, None))
        total = len(hits)
        shown = hits if not limit else hits[:limit]
        if hfilter is not None:
            for hit in shown:
                if hit["level"] is None:
                    ik = hit["id"].lower()
                    hit["level"] = self._node_level(ik, hfilter) if ik in self.nodes else None
        return ok(
            {
                "hits": shown,
                "truncated": len(shown) < total,
                "returned": len(shown),
                "total": total,
            }
        )

    def changelog_list(self, *, limit=20, session=False):
        entries = self.changelog[self.changelog_at_open :] if session else list(self.changelog)
        total = len(entries)
        shown = entries if not limit else entries[-limit:]
        out = []
        for ch in shown:
            src = ch.rows[-1] if ch.has_summary else ch.rows[0]
            out.append(
                {
                    "at": ch.at,
                    "type": ch.label,
                    "what": src.what,
                    "old": src.old,
                    "new": src.new,
                    "origin": ch.origin,
                    "n": ch.n,
                    "rows": [{"type": r.display_type, "what": r.what, "old": r.old, "new": r.new} for r in ch.rows],
                }
            )
        return ok(
            {
                "entries": out,
                "truncated": len(shown) < total,
                "returned": len(shown),
                "total": total,
            }
        )

    def warnings_list(self):
        return ok({"warnings": list(self.warnings)})

    def snapshot_sheet(self, type_="full sheet"):
        self.vs.append(
            {
                "type": type_,
                "og_file": None,
                "og_sheet": None,
                "build_warnings": list(self.warnings),
                "sheet": zlib.compress(pickle.dumps(self.data)),
                "required_data": self.snapshot_required_data(),
            }
        )

    def snapshot_ctrl_x_v_del_key_id_par(self):
        self.vs.append(
            {
                "type": "ctrl x, v, del key id par",
                "sheet": zlib.compress(pickle.dumps(self.data)),
                "required_data": self.snapshot_required_data(),
            }
        )

    def snapshot_ctrl_x_v_del_key(self):
        self.vs.append(
            {
                "type": "ctrl x, v, del key",
                "cells": {},
                "required_data": self.snapshot_required_data(),
            }
        )

    def associate_after_edit(self):
        first_hier = self.hiers[0]
        quick_hiers = self.hiers[1:]
        lh = len(self.hiers)
        to_insert = []
        if self.auto_sort_nodes_bool:
            for n in self.nodes.values():
                if all(p is None for p in n.ps.values()):
                    n.ps = {h: "" if n.cn[h] else None for h in self.hiers}
                    newrow = list(repeat("", self.row_len))
                    newrow[self.ic] = n.name
                    to_insert.append(newrow)
                tlly = 0
                for k, v in n.cn.items():
                    if v:
                        n.cn[k] = self.sort_node_cn(v, k)
                    elif not n.ps[k]:
                        n.ps[k] = None
                        tlly += 1
                if tlly == lh:
                    n.ps[first_hier] = ""
                    for h in quick_hiers:
                        n.ps[h] = None
        else:
            for n in self.nodes.values():
                if all(p is None for p in n.ps.values()):
                    n.ps = {h: "" if n.cn[h] else None for h in self.hiers}
                    newrow = list(repeat("", self.row_len))
                    newrow[self.ic] = n.name
                    to_insert.append(newrow)
                tlly = 0
                for k, v in n.cn.items():
                    if not v and not n.ps[k]:
                        n.ps[k] = None
                        tlly += 1
                if tlly == lh:
                    n.ps[first_hier] = ""
                    for h in quick_hiers:
                        n.ps[h] = None
        if to_insert:
            self.ops.insert_rows(to_insert)

    def rebuild_identity(self):
        self.auto_sort_nodes_bool = True
        built, nodes = TreeBuilder().build(
            input_sheet=self.data,
            output_sheet=[],
            row_len=self.row_len,
            ic=self.ic,
            hiers=self.hiers,
            nodes={},
            add_warnings=False,
            strip=not self.allow_spaces_ids_var,
        )
        self.data[:] = built
        self.nodes = nodes
        self.associate_after_edit()
        self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}

    def replace_mapping(self, mapping):
        mapping = {str(k).lower(): ("" if v is None else str(v)) for k, v in mapping.items()}
        idcols = set(self.hiers) | {self.ic}
        changes = []
        need_rebuild = False
        for r, row in enumerate(self.data):
            for c, cell in enumerate(row):
                new = cell
                for find, repl in mapping.items():
                    new = case_insensitive_replace(find, repl, new)
                if new != cell:
                    changes.append((r, c, cell, new))
                    if c in idcols:
                        need_rebuild = True
        if not changes:
            return ok({"file": None, "cells_changed": 0})
        if need_rebuild:
            self.snapshot_ctrl_x_v_del_key_id_par()
        else:
            self.snapshot_ctrl_x_v_del_key()
            for r, c, old, _new in changes:
                self.vs[-1]["cells"][(r, c)] = old
        log = ChangeBuilder()
        for r, c, old, new in changes:
            log.add(
                "Edit cell",
                f"ID: {self.data[r][self.ic]} column #{c + 1} named: {self.headers[c].name} with type: {self.headers[c].type_}",
                f"{old}",
                new,
            )
            self.data[r][c] = new
        cells_changed = len(changes)
        if cells_changed > 1:
            log.summary(f"Edit {cells_changed} cells")
        self.commit_change(log)
        if need_rebuild:
            self.rebuild_identity()
        else:
            self.rns = {row[self.ic].lower(): i for i, row in enumerate(self.data)}
        return ok({"file": None, "cells_changed": cells_changed})

    def _write_table(self, path, rows, *, overwrite=False):
        path = os.path.abspath(path)
        if os.path.exists(path) and not overwrite:
            return fail("file_exists", "File already exists   ")
        suffix = os.path.splitext(path)[1].lower()
        if suffix in (".xls", ".xlsm") or suffix not in (".csv", ".tsv", ".json", ".xlsx"):
            return fail("invalid_format", "Can only write .json, .xlsx or .csv    ")
        if suffix in (".csv", ".tsv"):
            with open(path, "w", newline="", encoding="utf-8") as fh:
                writer = csv.writer(
                    fh,
                    dialect=csv.excel_tab if suffix == ".tsv" else csv.excel,
                    lineterminator="\n",
                )
                writer.writerows(rows)
        elif suffix == ".json":
            with open(path, "w") as fh:
                fh.write(json.dumps({"records": [list(r) for r in rows]}, indent=4))
        else:
            wb = Workbook(write_only=True)
            ws = wb.create_sheet(title="Sheet1")
            for row in rows:
                ws.append(list(row))
            wb.active = wb["Sheet1"]
            wb.save(path)
        return ok({"file": path})

    def export_changelog(self, path, *, session=False, overwrite=False):
        changes = self.changelog[self.changelog_at_open :] if session else self.changelog
        rows, _ = flatten_changelog(changes)
        out = self._write_table(path, [list(r) for r in rows], overwrite=overwrite)
        if not out["ok"]:
            return out
        return ok({"file": out["result"]["file"], "rows": len(rows)})

    def export_flat(
        self,
        path,
        *,
        hier,
        details=True,
        justify=True,
        reverse=False,
        index=False,
        remove_end_ids=0,
        overwrite=False,
    ):
        if isinstance(hier, int):
            if hier not in self.hiers:
                return fail("unknown_hierarchy", f"Unknown hierarchy {hier}   ")
            h = hier
        else:
            hout = self.resolve_hier(hier)
            if not hout["ok"]:
                return hout
            h = hout["result"]["index"]
        flat = TreeBuilder().build_flattened(
            input_sheet=self.data,
            output_sheet=[],
            nodes=self.nodes,
            headers=[f"{hdr.name}" for hdr in self.headers],
            ic=int(self.ic),
            pc=int(h),
            hiers=list(self.hiers),
            detail_columns=details,
            justify_left=justify,
            reverse=reverse,
            add_index=index,
            remove_end_ids=remove_end_ids,
        )
        out = self._write_table(path, flat, overwrite=overwrite)
        if not out["ok"]:
            return out
        cols = len(flat[0]) if flat else 0
        return ok({"file": out["result"]["file"], "rows": len(flat), "cols": cols})

    def restore_snapshot(self, entry):
        """Restore sheet/nodes from a vs entry. Must not touch changelog."""
        rd = entry["required_data"]
        self.ic = rd["ic"]
        self.pc = rd["pc"]
        self.hiers = rd["hiers"]
        self.nodes = pickle.loads(rd["nodes"])
        self.row_len = rd["row_len"]
        self.auto_sort_nodes_bool = rd["auto_sort_nodes_bool"]
        self.topnodes_order = rd["topnodes_order"]
        self.tagged_ids = rd["tagged_ids"]
        self.headers = rd["headers"]
        typ = entry["type"]
        if typ == "add id":
            rn = entry["row"]["rn"]
            if entry["row"].get("added_or_changed") == "changed":
                self.data[rn] = entry["row"]["stored"]
            elif entry["row"].get("added_or_changed") == "added":
                del self.data[rn]
        elif typ == "rename id":
            for tup in entry["rows"]:
                rn, h, v = pickle.loads(zlib.decompress(tup))
                self.data[rn][h] = v
            self.data[entry["ikrow"][0]][self.ic] = entry["ikrow"][2]
        elif typ == "paste id":
            for tup in entry["rows"]:
                rn, fromcol, frompar, tocol, topar = pickle.loads(zlib.decompress(tup))
                self.data[rn][fromcol] = frompar
                self.data[rn][tocol] = topar
        elif typ == "delete ids":
            rows = entry["rows"]
            for rn in sorted(r for r, obj in rows.items() if obj.t == 1):
                self.data.insert(rn, rows[rn].row)
            for rn in sorted(r for r, obj in rows.items() if obj.t == 0):
                for h, par in zip(self.hiers, pickle.loads(zlib.decompress(rows[rn].row))):
                    self.data[rn][h] = par
        elif typ == "add col":
            c = entry["treecolsel"]
            for r in range(len(self.data)):
                del self.data[r][c]
        elif typ == "del cols":
            for cn, rowdict in reversed(entry["cols"].items()):
                for rn, v in rowdict.items():
                    self.data[rn].insert(cn, v)
        elif typ == "edit validation":
            for rn, c in enumerate(pickle.loads(zlib.decompress(entry["col"]))):
                self.data[rn][entry["col_num"]] = c
        elif typ == "sort":
            self.data = [self.data[self.rns[entry["ids"][oldrn]]] for oldrn in range(len(entry["ids"]))]
        elif typ == "prune changelog":
            pass
        elif typ == "drag rows":
            mapping = dict(zip(entry["row_mapping"].values(), entry["row_mapping"]))
            self.data[:] = _move_elements(self.data, mapping)
        elif typ == "drag cols":
            mapping = entry["column_mapping"]
            self.data[:] = [_move_elements(row, mapping) for row in self.data]
        elif typ.startswith("full"):
            self.warnings = entry.get("build_warnings", self.warnings)
            self.data = pickle.loads(zlib.decompress(entry["sheet"]))
        elif typ == "ctrl x, v, del key id par":
            self.data = pickle.loads(zlib.decompress(entry["sheet"]))
        elif typ == "ctrl x, v, del key":
            for k, v in entry["cells"].items():
                self.data[k[0]][k[1]] = v
        self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}

    def undo(self):
        if not self.vs:
            return fail("nothing_to_undo", "Nothing to undo   ")
        self._pending = None
        entry = self.vs.pop()
        typ = entry["type"]
        self.restore_snapshot(entry)
        if self.changelog:
            self.changelog.pop()
        if typ == "prune changelog":
            self.changelog = entry["rows"] + self.changelog
            self.changelog_at_open = entry.get("changelog_at_open", self.changelog_at_open)
        return ok({"undone": typ, "undo": len(self.vs)})

    def _convert_merge_rows(self, rows, fmt, id_col, parent_cols):
        incoming = [list(r) for r in rows]
        warnings = []
        row_len = equalize_sublist_lens(incoming)
        if fmt == 0:
            if id_col is None or not parent_cols:
                return fail("need_columns", "Need ID and parent columns   ")
            ns_ic = id_col
            ns_hiers = list(parent_cols)
        elif fmt in (1, 2, 3, 4):
            if not parent_cols:
                return fail("need_columns", "Need parent columns   ")
            incoming, row_len, ns_ic, ns_hiers = TreeBuilder().convert_flattened_to_normal(
                data=incoming,
                hier_cols=parent_cols,
                rowlen=row_len,
                fmt=fmt,
                warnings=warnings,
            )
        elif fmt == 5:
            incoming, row_len, ns_ic, ns_hiers = TreeBuilder().convert_indented_tree_detail_adjacent_to_normal(
                data=incoming,
            )
        elif fmt == 6:
            incoming, row_len, ns_ic, ns_hiers = TreeBuilder().convert_indented_tree_details_adjacent_to_normal(
                data=incoming,
            )
        elif fmt == 7:
            incoming, row_len, ns_ic, ns_hiers = TreeBuilder().convert_indented_tree_with_header_to_normal(
                data=incoming,
            )
        else:
            return fail("invalid_format", "Unknown format   ")
        if not incoming:
            return fail("invalid_format", "File contains no data   ")
        ns_headers = self.fix_headers(incoming.pop(0), row_len, warnings=False)
        equalize_sublist_lens(seq=incoming, len_=len(ns_headers))
        return ok(
            {
                "incoming": incoming,
                "ns_ic": ns_ic,
                "ns_hiers": ns_hiers,
                "ns_headers": ns_headers,
                "warnings": warnings,
            }
        )

    def merge_from_rows(
        self,
        rows,
        *,
        fmt=0,
        id_col=None,
        parent_cols=None,
        add_ids=True,
        add_dcols=True,
        add_pcols=True,
        overwrite_details=True,
        overwrite_parents=True,
        insert_row=None,
        file_opened="",
    ):
        converted = self._convert_merge_rows(rows, fmt, id_col, parent_cols)
        if not converted["ok"]:
            return converted
        incoming = converted["result"]["incoming"]
        ns_ic = converted["result"]["ns_ic"]
        ns_hiers = converted["result"]["ns_hiers"]
        ns_headers = converted["result"]["ns_headers"]
        self.warnings = list(converted["result"]["warnings"])
        self._ensure_snapshot("full sheet", self.snapshot_sheet)
        log = ChangeBuilder(origin=ORIGIN_MERGE)
        ns_hiers_set = set(ns_hiers)
        ns_pcol_names = {cell.lower(): i for i, cell in enumerate(ns_headers) if i in ns_hiers_set}
        ns_dcol_names = {cell.lower(): i for i, cell in enumerate(ns_headers) if i not in ns_hiers_set and i != ns_ic}
        os_header_names = {h.name.lower() for h in self.headers}
        ns_rns = {row[ns_ic].lower(): i for i, row in enumerate(incoming)}
        shared_ids = {i: ik for ik, i in self.rns.items() if ik in ns_rns}
        os_pcol_names = {h.name.lower(): i for i, h in enumerate(self.headers) if h.type_ == "Parent"}
        os_dcol_names = {h.name.lower(): i for i, h in enumerate(self.headers) if h.type_ == "Text"}
        changes_made = 0
        ids_added = 0
        details_written = 0
        parents_written = 0
        detail_columns_added = 0
        parent_columns_added = 0
        rows_to_insert = []

        if add_dcols:
            new_dcols = [idx for colname, idx in ns_dcol_names.items() if colname not in os_header_names]
            num_new_dcols = len(new_dcols)
            self.headers.extend([Header(ns_headers[idx], "Text") for idx in new_dcols])
            if num_new_dcols:
                self.ops.insert_cols(self.row_len, num_new_dcols)
            for num, idx in enumerate(new_dcols, 1):
                self._log_prefixed(
                    log,
                    "Merge | Add new detail column",
                    f"Column #{self.row_len + num} with name: {ns_headers[idx]} and type: Text",
                    "",
                    "",
                )
                changes_made += 1
                detail_columns_added += 1
            for rn in range(len(self.data)):
                row = self.data[rn]
                if rn in shared_ids:
                    ns_rn = ns_rns[shared_ids[rn]]
                    for num, idx in enumerate(new_dcols):
                        row[self.row_len + num] = incoming[ns_rn][idx]
                        if row[self.row_len + num] != "":
                            self._log_prefixed(
                                log,
                                "Merge | Edit cell",
                                f"ID: {row[self.ic]} column #{self.row_len + num + 1} named: {self.headers[self.row_len + num].name} with type: {self.headers[self.row_len + num].type_}",
                                "",
                                f"{row[self.row_len + num]}",
                            )
                            changes_made += 1
                            details_written += 1
                self.data[rn] = row
            self.row_len += num_new_dcols

        if add_pcols:
            new_pcols = [idx for colname, idx in ns_pcol_names.items() if colname not in os_header_names]
            num_new_pcols = len(new_pcols)
            self.headers.extend([Header(ns_headers[idx], "Parent") for idx in new_pcols])
            if num_new_pcols:
                self.ops.insert_cols(self.row_len, num_new_pcols)
            for num, idx in enumerate(new_pcols, 1):
                self._log_prefixed(
                    log,
                    "Merge | Add new hierarchy column",
                    f"Column #{self.row_len + num} with name: {ns_headers[idx]}",
                    "",
                    "",
                )
                changes_made += 1
                parent_columns_added += 1
            range_end = self.row_len + num_new_pcols
            self.hiers.extend(list(range(self.row_len, range_end)))
            for node in self.nodes.values():
                for i in range(self.row_len, range_end):
                    node.ps[i] = None
                    node.cn[i] = []
            for rn in range(len(self.data)):
                row = self.data[rn]
                if rn in shared_ids:
                    ns_rn = ns_rns[shared_ids[rn]]
                    for num, idx in enumerate(new_pcols):
                        row[self.row_len + num] = incoming[ns_rn][idx]
                        if row[self.row_len + num] != "":
                            self._log_prefixed(
                                log,
                                "Merge | Edit cell",
                                f"ID: {row[self.ic]} column #{self.row_len + num + 1} named: {self.headers[self.row_len + num].name} with type: {self.headers[self.row_len + num].type_}",
                                "",
                                f"{row[self.row_len + num]}",
                            )
                            changes_made += 1
                            parents_written += 1
                self.data[rn] = row
            self.row_len += num_new_pcols

        if add_ids:
            new_ids = {ik for ik in ns_rns if ik not in self.rns and ik}
            shared_dcols = tuple(name for name in os_dcol_names if name in ns_dcol_names)
            shared_pcols = tuple(name for name in os_pcol_names if name in ns_pcol_names)
            new_dcol_indexes = {
                i: h.name.lower()
                for i, h in enumerate(self.headers)
                if add_dcols and h.name.lower() in ns_dcol_names and h.name.lower() not in os_dcol_names
            }
            new_pcol_indexes = {
                i: h.name.lower()
                for i, h in enumerate(self.headers)
                if add_pcols and h.name.lower() in ns_pcol_names and h.name.lower() not in os_pcol_names
            }
            for row in incoming:
                if row[ns_ic].lower() not in new_ids:
                    continue
                newrow = list(repeat("", self.row_len))
                newrow[self.ic] = row[ns_ic]
                self._log_prefixed(
                    log,
                    "Merge | Add ID",
                    f"Name: {newrow[self.ic]} Parent: n/a - Top ID column #{self.hiers[0] + 1} named: {self.headers[self.hiers[0]].name}",
                    "",
                    "",
                )
                changes_made += 1
                ids_added += 1
                for idx, colname in new_dcol_indexes.items():
                    newrow[idx] = row[ns_dcol_names[colname]]
                    if newrow[idx] != "":
                        self._log_prefixed(
                            log,
                            "Merge | Edit cell",
                            f"ID: {newrow[self.ic]} column #{idx + 1} named: {self.headers[idx].name} with type: {self.headers[idx].type_}",
                            "",
                            f"{newrow[idx]}",
                        )
                        changes_made += 1
                        details_written += 1
                for idx, colname in new_pcol_indexes.items():
                    newrow[idx] = row[ns_pcol_names[colname]]
                    if newrow[idx] != "":
                        self._log_prefixed(
                            log,
                            "Merge | Edit cell",
                            f"ID: {newrow[self.ic]} column #{idx + 1} named: {self.headers[idx].name} with type: {self.headers[idx].type_}",
                            "",
                            f"{newrow[idx]}",
                        )
                        changes_made += 1
                        parents_written += 1
                for name in shared_dcols:
                    if self.detail_is_valid_for_col(os_dcol_names[name], row[ns_dcol_names[name]]):
                        newrow[os_dcol_names[name]] = row[ns_dcol_names[name]]
                        hdr_idx = os_dcol_names[name]
                        if newrow[hdr_idx] != "":
                            self._log_prefixed(
                                log,
                                "Merge | Edit cell",
                                f"ID: {newrow[self.ic]} column #{hdr_idx + 1} named: {self.headers[hdr_idx].name} with type: {self.headers[hdr_idx].type_}",
                                "",
                                f"{newrow[hdr_idx]}",
                            )
                            changes_made += 1
                            details_written += 1
                for name in shared_pcols:
                    newrow[os_pcol_names[name]] = row[ns_pcol_names[name]]
                    hdr_idx = os_pcol_names[name]
                    if newrow[hdr_idx] != "":
                        self._log_prefixed(
                            log,
                            "Merge | Edit cell",
                            f"ID: {newrow[self.ic]} column #{hdr_idx + 1} named: {self.headers[hdr_idx].name} with type: {self.headers[hdr_idx].type_}",
                            "",
                            f"{newrow[hdr_idx]}",
                        )
                        changes_made += 1
                        parents_written += 1
                rows_to_insert.append(newrow)

        if overwrite_details:
            shared_dcols = {name: idx for name, idx in os_dcol_names.items() if name in ns_dcol_names}
            for rn in range(len(self.data)):
                row = self.data[rn]
                if rn in shared_ids:
                    ns_rn = ns_rns[shared_ids[rn]]
                    for name, idx in shared_dcols.items():
                        ns_dcol_idx = ns_dcol_names[name]
                        if (
                            self.detail_is_valid_for_col(idx, incoming[ns_rn][ns_dcol_idx])
                            and row[idx] != incoming[ns_rn][ns_dcol_idx]
                        ):
                            self._log_prefixed(
                                log,
                                "Merge | Edit cell",
                                f"ID: {row[self.ic]} column #{idx + 1} named: {self.headers[idx].name} with type: {self.headers[idx].type_}",
                                f"{row[idx]}",
                                incoming[ns_rn][ns_dcol_idx],
                            )
                            changes_made += 1
                            details_written += 1
                            row[idx] = incoming[ns_rn][ns_dcol_idx]
                self.data[rn] = row

        if overwrite_parents:
            shared_pcols = {name: idx for name, idx in os_pcol_names.items() if name in ns_pcol_names}
            for rn in range(len(self.data)):
                row = self.data[rn]
                if rn in shared_ids:
                    ns_rn = ns_rns[shared_ids[rn]]
                    for name, idx in shared_pcols.items():
                        ns_pcol_idx = ns_pcol_names[name]
                        if row[idx] != incoming[ns_rn][ns_pcol_idx]:
                            self._log_prefixed(
                                log,
                                "Merge | Edit cell",
                                f"ID: {row[self.ic]} column #{idx + 1} named: {self.headers[idx].name} with type: {self.headers[idx].type_}",
                                f"{row[idx]}",
                                incoming[ns_rn][ns_pcol_idx],
                            )
                            changes_made += 1
                            parents_written += 1
                            row[idx] = incoming[ns_rn][ns_pcol_idx]
                self.data[rn] = row

        if rows_to_insert:
            self.ops.insert_rows(rows_to_insert, insert_row)
        if changes_made:
            log.summary(
                f"Merged sheets making {changes_made} {'changes' if changes_made > 1 else 'change'}",
                f"{'With file:' if file_opened else ''} {file_opened}",
            )
            self.commit_change(log)
            self.nodes = {}
            self.auto_sort_nodes_bool = True
            built, nodes, warnings = TreeBuilder().build(
                self.data,
                [],
                self.row_len,
                self.ic,
                self.hiers,
                self.nodes,
                warnings=self.warnings,
                add_warnings=True,
                strip=not self.allow_spaces_ids_var,
            )
            self.nodes = nodes
            self.warnings = warnings
            self.data[:] = built
            self.associate(startup=False)
            self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
            return ok(
                {
                    "ids_added": ids_added,
                    "details_written": details_written,
                    "parents_written": parents_written,
                    "detail_columns_added": detail_columns_added,
                    "parent_columns_added": parent_columns_added,
                },
                warnings=list(self.warnings),
            )
        self.vs.pop()
        return fail("no_changes", "No applicable changes were made")

    def cut_paste_children(self, oldparent, newparent, hier, snapshot=True):
        self.refresh_rows = set()
        pk = oldparent.lower()
        npk = newparent.lower()
        if pk not in self.nodes:
            return fail("id_not_found", "ID doesn't exist   ")
        if not len(self.nodes[pk].cn[hier]):
            return fail("id_not_found", f"{self.nodes[pk].name} has no children   ")
        already_in = set()
        if hier != self.pc:
            for ciid in self.nodes[pk].cn[hier]:
                for diid in self.check_cn(ciid, hier):
                    if self.nodes[diid].ps[self.pc] is not None:
                        already_in.add(ciid)
                        break
            if len(already_in) == len(self.nodes[pk].cn[hier]):
                return fail(
                    "already_in_hierarchy",
                    f"Unable to move children, key IDs are already in {self.headers[self.pc].name}   ",
                )
        else:
            if any(npk == ck for ck in self.check_cn(pk, hier)):
                return fail("cycle", "Cannot add ID to same line   ")
            if pk == npk:
                return fail("already_in_hierarchy", "Children already have this parent   ")
        if snapshot:
            self._ensure_snapshot("paste id", self.snapshot_paste_id)
        for ciid in tuple(self.nodes[pk].cn[hier]):
            if ciid in already_in:
                continue
            if not self.auto_sort_nodes_bool and npk == "":
                self.topnodes_order[self.pc].append(ciid)
            crow = self.rns[ciid]
            if snapshot:
                self.refresh_rows.add(int(crow))
                self.vs[-1]["rows"].append(
                    zlib.compress(
                        pickle.dumps(
                            (
                                crow,
                                hier,
                                self.data[crow][hier],
                                self.pc,
                                self.data[crow][self.pc],
                            )
                        )
                    )
                )
            self.data[crow][hier] = ""
            self.nodes[ciid].ps[hier] = None
            if npk:
                self.data[crow][self.pc] = self.nodes[npk].name
                self.nodes[ciid].ps[self.pc] = npk
                self.nodes[npk].cn[self.pc].append(ciid)
            else:
                self.data[crow][self.pc] = ""
                self.nodes[ciid].ps[self.pc] = ""
            self.nodes[pk].cn[hier].remove(ciid)
            if hier != self.pc:
                stack = [ciid]
                while stack:
                    current_iid = stack.pop()
                    children = [
                        child_iid for child_iid in self.nodes[current_iid].cn[hier] if child_iid not in already_in
                    ]
                    self.nodes[current_iid].cn[self.pc] = children
                    children_to_remove = set(children)
                    self.nodes[current_iid].cn[hier] = [
                        child for child in self.nodes[current_iid].cn[hier] if child not in children_to_remove
                    ]
                    for child_iid in reversed(children):
                        child = self.nodes[child_iid]
                        crow = self.rns[child_iid]
                        if snapshot:
                            self.refresh_rows.add(int(crow))
                            self.vs[-1]["rows"].append(
                                zlib.compress(
                                    pickle.dumps(
                                        (
                                            crow,
                                            hier,
                                            self.data[crow][hier],
                                            self.pc,
                                            self.data[crow][self.pc],
                                        )
                                    )
                                )
                            )
                        self.data[crow][self.pc] = f"{self.data[crow][hier]}"
                        self.data[crow][hier] = ""
                        child.ps[self.pc] = current_iid
                        child.ps[hier] = None
                        stack.append(child_iid)
        if self.auto_sort_nodes_bool:
            if self.nodes[pk].ps[hier]:
                parent_parent_node = self.nodes[self.nodes[pk].ps[hier]]
                parent_parent_node.cn[hier] = self.sort_node_cn(parent_parent_node.cn[hier], hier)
            if npk:
                if self.nodes[npk].ps[self.pc]:
                    parent_parent_node = self.nodes[self.nodes[npk].ps[self.pc]]
                    parent_parent_node.cn[self.pc] = self.sort_node_cn(parent_parent_node.cn[self.pc], self.pc)
                self.nodes[npk].cn[self.pc] = self.sort_node_cn(self.nodes[npk].cn[self.pc], self.pc)
        return ok({})

    def _drop_iids(self, to_del):
        if to_del:
            self.ops.delete_rows([self.rns[iid] for iid in to_del if iid in self.rns])
        self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}

    def _import_one_change(self, change, log):
        ctyp = change[1]
        if ctyp.startswith("Imported change |"):
            ctyp = ctyp.split("Imported change | ")[1]
        elif ctyp.startswith("Merge |"):
            ctyp = ctyp.split("Merge | ")[1]
        if ctyp == "Edit cell |" or ctyp == "Edit cell":
            c3s = change[2].split(" ")
            cik = c3s[1].lower()
            name = c3s[5]
            try:
                col = next(i for i, h in enumerate(self.headers) if h.name.lower() == name.lower())
            except StopIteration:
                return False, "column_missing"
            type_ = c3s[-1]
            if type_ == "Detail":
                type_ = f"{c3s[-2]} {type_}"
            if self.headers[col].validation:
                validation_check = self.is_in_validation(self.headers[col].validation, change[4])
            else:
                validation_check = True
            if self.headers[col].type_ != normalize_header_type(type_):
                return False, "type_mismatch"
            if cik not in self.rns:
                return False, "id_missing"
            if self.data[self.rns[cik]][col] != change[3]:
                return False, "value_mismatch"
            if not validation_check:
                return False, "validation"
            if self.data[self.rns[cik]][col] != change[4]:
                self._log_prefixed(log, "Imported change | Edit cell", change[2], change[3], change[4])
                self.data[self.rns[cik]][col] = change[4]
                if type_ == "ID" or type_ == "Parent":
                    self.rebuild_identity()
                return True, None
            return True, "unnecessary"
        if ctyp == "Move rows":
            old_locs = change[3].split(",")
            new_locs = change[4].split(",")
            if len(old_locs) != len(new_locs):
                return False, "other"
            if len(old_locs) == 1:
                old_locs = [old_locs[0].split("Old locations: ")[1]]
                new_locs = [new_locs[0].split("New locations: ")[1]]
            new_idxs = dict(zip(map(int, old_locs), map(int, new_locs)))
            if all(0 <= i <= len(self.data) for i in new_idxs) and all(
                0 <= i <= len(self.data) for i in new_idxs.values()
            ):
                self.data[:] = _move_elements(self.data, new_idxs)
                self._log_prefixed(log, "Imported change | Move rows", change[2], change[3], change[4])
                self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
                return True, None
            return False, "other"
        if ctyp == "Move columns":
            old_locs = change[3].split(",")
            new_locs = change[4].split(",")
            if len(old_locs) != len(new_locs):
                return False, "other"
            if len(old_locs) == 1:
                old_locs = [old_locs[0].split("Old locations: ")[1]]
                new_locs = [new_locs[0].split("New locations: ")[1]]
            new_idxs = dict(zip(map(int, old_locs), map(int, new_locs)))
            if max(new_idxs.values()) < self.row_len:
                full = _full_new_idxs(self.row_len, new_idxs)
                self.data[:] = [_move_elements(row, new_idxs) for row in self.data]
                self.ic = full[self.ic]
                self.pc = full[self.pc]
                self.headers = _move_elements(self.headers, new_idxs)
                self.hiers = sorted(full[c] for c in self.hiers)
                for node in self.nodes.values():
                    node.cn = {full[k]: v for k, v in node.cn.items()}
                    node.ps = {full[k]: v for k, v in node.ps.items()}
                if not self.auto_sort_nodes_bool:
                    self.topnodes_order = {full[k]: v for k, v in self.topnodes_order.items()}
                self._log_prefixed(log, "Imported change | Move columns", change[2], change[3], change[4])
                return True, None
            return False, "other"
        if ctyp == "Add new hierarchy column":
            c3s = change[2].split(" ")
            colname = "".join(c3s[-1].split(" ")).strip()
            colnum = int(c3s[1][1:]) - 1
            if colname.lower() not in (h.name.lower() for h in self.headers) and 0 <= colnum <= len(self.headers):
                self.add_hier_col(colnum, colname, snapshot=False)
                self._log_prefixed(
                    log,
                    "Imported change | Add new hierarchy column",
                    change[2],
                    change[3],
                    change[4],
                )
                return True, None
            return False, "other"
        if ctyp == "Add new detail column":
            c3s = change[2].split(" ")
            colname = "".join(c3s[4].split(" ")).strip()
            colnum = int(c3s[1][1:]) - 1
            coltype = f"{c3s[-2]} {c3s[-1]}"
            if colname.lower() not in (h.name.lower() for h in self.headers) and 0 <= colnum <= len(self.headers):
                self.add_col(colnum, colname, coltype, snapshot=False)
                self._log_prefixed(
                    log,
                    "Imported change | Add new detail column",
                    change[2],
                    change[3],
                    change[4],
                )
                return True, None
            return False, "other"
        if ctyp == "Delete hierarchy column":
            c3s = change[2].split(" ")
            colname = c3s[-1]
            try:
                colnum = next(i for i, h in enumerate(self.headers) if h.name.lower() == colname.lower())
            except StopIteration:
                return False, "column_missing"
            if self.headers[colnum].type_ == "Parent" and len(self.hiers) > 1:
                if self.pc == colnum:
                    self.pc = int(next(i for i in self.hiers if i != colnum))
                self.del_cols(cols=[colnum], snapshot=False)
                self._log_prefixed(log, "Imported change | Delete hierarchy column", change[2], "", "")
                return True, None
            return False, "other"
        if ctyp == "Delete detail column":
            c3s = change[2].split(" ")
            colname = c3s[4]
            try:
                colnum = next(i for i, h in enumerate(self.headers) if h.name.lower() == colname.lower())
            except StopIteration:
                return False, "column_missing"
            coltype = f"{c3s[-2]} {c3s[-1]}"
            if self.headers[colnum].type_ == "Text" and normalize_header_type(coltype) == "Text":
                self.del_cols(cols=[colnum], snapshot=False)
                self._log_prefixed(log, "Imported change | Delete detail column", change[2], "", "")
                return True, None
            return False, "other"
        if ctyp == "Column rename":
            c3s = change[2].split(" ")
            coltype = f"{c3s[-2]} {c3s[-1]}"
            colname = "".join(change[4].split(" ")).strip()
            try:
                colnum = next(i for i, h in enumerate(self.headers) if h.name.lower() == colname.lower())
            except StopIteration:
                return False, "column_missing"
            if (
                self.headers[colnum].name.lower() == change[3].lower()
                and self.headers[colnum].type_ == normalize_header_type(coltype)
                and colname.lower() not in (h.name.lower() for h in self.headers)
            ):
                self.rename_col(colnum, colname, snapshot=False)
                self._log_prefixed(log, "Imported change | Column rename", change[2], change[3], change[4])
                return True, None
            return False, "other"
        if ctyp == "Edit validation":
            c3s = change[2].split(" ")
            colname = c3s[3]
            try:
                colnum = next(i for i, h in enumerate(self.headers) if h.name.lower() == colname.lower())
            except StopIteration:
                return False, "column_missing"
            coltype = f"{c3s[-2]} {c3s[-1]}"
            validation = change[4]
            if (
                self.headers[colnum].type_ == "Text"
                and normalize_header_type(coltype) == "Text"
                and change[3] == ",".join(self.headers[colnum].validation)
            ):
                if validation:
                    validation = self.check_validation_validity(colnum, validation.split(","))
                    if isinstance(validation, str):
                        return False, "validation"
                else:
                    validation = []
                self.headers[colnum].validation = validation
                if validation:
                    self.apply_validation_to_col(colnum)
                self._log_prefixed(log, "Imported change | Edit validation", change[2], change[3], change[4])
                return True, None
            return False, "other"
        if ctyp in ("Date format change", "Change detail column type"):
            return False, "other"
        if ctyp in ("Cut and paste ID", "Cut and paste ID |"):
            return self._import_cut_or_copy(change, log, all_children=False, copy=False)
        if ctyp in ("Cut and paste ID + children", "Cut and paste ID + children |"):
            return self._import_cut_or_copy(change, log, all_children=True, copy=False)
        if ctyp == "Cut and paste children":
            return self._import_cut_children(change, log)
        if ctyp in ("Copy and paste ID |", "Copy and paste ID"):
            return self._import_cut_or_copy(change, log, all_children=False, copy=True)
        if ctyp in ("Copy and paste ID + children |", "Copy and paste ID + children"):
            return self._import_cut_or_copy(change, log, all_children=True, copy=True)
        if ctyp == "Add ID":
            new = change[2].split(" ")
            newcolname = new[-1]
            try:
                newcol = next(i for i, h in enumerate(self.headers) if h.name.lower() == newcolname.lower())
            except StopIteration:
                return False, "column_missing"
            cid = new[1]
            if "n/a - Top ID" in change[2]:
                newpar = ""
                newpk = ""
            else:
                newpar = new[3]
                newpk = newpar.lower()
            newpar_check = bool(not newpk or newpk in self.rns)
            if self.headers[newcol].type_ == "Parent" and newpar_check:
                oldpc = int(self.pc)
                self.pc = newcol
                out = self.add(cid, newpar, snapshot=False)
                self.pc = int(oldpc)
                if out["ok"]:
                    self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
                    self._log_prefixed(log, "Imported change | Add ID", change[2], change[3], change[4])
                    return True, None
                return False, "parent_mismatch"
            return False, "parent_mismatch"
        if ctyp == "Rename ID":
            oldname = change[3]
            newname = change[4]
            if oldname.lower() in self.rns and newname.lower() not in self.rns:
                out = self.rename(oldname, newname, snapshot=False)
                if out["ok"]:
                    self._log_prefixed(log, "Imported change | Rename ID", change[2], change[3], change[4])
                    self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
                    return True, None
                return False, "other"
            return False, "id_missing"
        if ctyp in ("Delete ID |", "Delete ID"):
            return self._import_delete(change, log, children=False, all_hierarchies=False, orphan=False)
        if ctyp == "Delete ID, orphan children":
            return self._import_delete(change, log, children=False, all_hierarchies=False, orphan=True)
        if ctyp in ("Delete ID + all children |", "Delete ID + all children"):
            return self._import_delete(change, log, children=True, all_hierarchies=False, orphan=False)
        if ctyp in (
            "Delete ID + all children from all hierarchies |",
            "Delete ID + all children from all hierarchies",
        ):
            return self._import_delete(change, log, children=True, all_hierarchies=True, orphan=False)
        if ctyp in ("Delete ID from all hierarchies |", "Delete ID from all hierarchies"):
            cid = change[2]
            if cid.lower() in self.rns:
                to_del = self._del_id_all_core(cid.lower(), snapshot=False)
                self._drop_iids(to_del)
                self._log_prefixed(
                    log,
                    "Imported change | Delete ID from all hierarchies",
                    change[2],
                    change[3],
                    change[4],
                )
                return True, None
            return False, "id_missing"
        if ctyp == "Delete ID from all hierarchies, orphan children":
            cid = change[2]
            if cid.lower() in self.rns:
                to_del = self._del_id_all_orphan_core(cid.lower(), snapshot=False)
                self._drop_iids(to_del)
                self._log_prefixed(
                    log,
                    "Imported change | Delete ID from all hierarchies, orphan children",
                    change[2],
                    change[3],
                    change[4],
                )
                return True, None
            return False, "id_missing"
        if ctyp == "Sort sheet":
            if change[2] == "Sorted sheet in tree walk order":
                if self.data:
                    self.sort_sheet_walk(snapshot=False)
                    self._log_prefixed(log, f"Imported change | {change[1]}", change[2], change[3], change[4])
                    return True, None
                return False, "other"
            c3s = change[2].split(" ")
            colname = c3s[6]
            order = c3s[8]
            if order in ("ASCENDING", "DESCENDING"):
                self.sort_sheet(colname, order, snapshot=False)
                self._log_prefixed(log, f"Imported change | {change[1]}", change[2], change[3], change[4])
                return True, None
            return False, "other"
        return False, "other"

    def _import_delete(self, change, log, *, children, all_hierarchies, orphan):
        info = change[2].split(" ")
        colname = info[-1]
        try:
            colnum = next(i for i, h in enumerate(self.headers) if h.name.lower() == colname.lower())
        except StopIteration:
            return False, "column_missing"
        cid = info[1]
        cpar = "" if "n/a - Top ID" in change[2] else info[3]
        if cpar:
            if cpar.lower() not in self.nodes or self.nodes[self.nodes[cid.lower()].ps[colnum]].name != cpar:
                cpar_check = False
            else:
                cpar_check = True
        else:
            cpar_check = True
        if cid.lower() in self.rns and cpar_check and self.headers[colnum].type_ == "Parent":
            oldpc = int(self.pc)
            self.pc = colnum
            if orphan:
                to_del = self._del_id_orphan_core(cid.lower(), cpar.lower(), snapshot=False)
                self._drop_iids(to_del)
            elif children and all_hierarchies:
                to_del = self._del_id_children_all_core(cid.lower(), snapshot=False)
                self._drop_iids(to_del)
            elif children:
                to_del = self._del_id_children_core(cid.lower(), snapshot=False)
                self._drop_iids(to_del)
            else:
                to_del = self._del_id_core(cid.lower(), snapshot=False)
                self._drop_iids(to_del)
            self.pc = int(oldpc)
            self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
            label = (
                change[1]
                if change[1].startswith("Imported change |")
                else f"Imported change | {change[1].rstrip(' |')}"
            )
            if orphan:
                label = "Imported change | Delete ID"
            elif children and all_hierarchies:
                label = "Imported change | Delete ID + all children from all hierarchies"
            elif children:
                label = "Imported change | Delete ID + all children"
            else:
                label = "Imported change | Delete ID"
            self._log_prefixed(log, label, change[2], change[3], change[4])
            return True, None
        if cid.lower() not in self.rns:
            return False, "id_missing"
        return False, "parent_mismatch"

    def _import_cut_or_copy(self, change, log, *, all_children, copy):
        cik = change[2].lower()
        old = change[3].split(" ")
        oldcolname = old[-1]
        try:
            oldcol = next(i for i, h in enumerate(self.headers) if h.name.lower() == oldcolname.lower())
            new = change[4].split(" ")
            newcolname = new[-1]
            newcol = next(i for i, h in enumerate(self.headers) if h.name.lower() == newcolname.lower())
        except StopIteration:
            return False, "column_missing"
        if copy:
            oldpar_check = True
        elif "n/a - Top ID" in change[3]:
            oldpar = ""
            oldpar_check = True
        else:
            oldpar = old[2]
            if oldpar.lower() not in self.nodes or oldpar != self.nodes[self.nodes[cik].ps[oldcol]].name:
                oldpar_check = False
            else:
                oldpar_check = True
        if "n/a - Top ID" in change[4]:
            newpar = ""
            newpar_check = True
        else:
            newpar = new[2]
            if newpar.lower() not in self.nodes or self.nodes[newpar.lower()].ps[newcol] is None:
                newpar_check = False
            else:
                newpar_check = True
        if (
            self.headers[oldcol].type_ == "Parent"
            and self.headers[newcol].type_ == "Parent"
            and cik in self.rns
            and oldpar_check
            and newpar_check
        ):
            oldpc = int(self.pc)
            self.pc = newcol
            if copy and all_children:
                out = self.copy_paste_all(change[2], oldcol, newpar, snapshot=False)
            elif copy:
                out = self.copy_paste(change[2], oldcol, newpar, snapshot=False)
            elif all_children:
                out = self.cut_paste_all(change[2], oldpar, oldcol, newpar, snapshot=False)
            else:
                out = self.cut_paste(change[2], oldpar, oldcol, newpar, snapshot=False)
            self.pc = int(oldpc)
            if out["ok"]:
                if copy and all_children:
                    label = "Imported change | Copy and paste ID + children"
                elif copy:
                    label = "Imported change | Copy and paste ID"
                elif all_children:
                    label = "Imported change | Cut and paste ID + children"
                else:
                    label = "Imported change | Cut and paste ID"
                self._log_prefixed(log, label, change[2], change[3], change[4])
                return True, None
            return False, "parent_mismatch"
        if cik not in self.rns:
            return False, "id_missing"
        return False, "parent_mismatch"

    def _import_cut_children(self, change, log):
        old = change[3].split(" ")
        oldcolname = old[-1]
        try:
            oldcol = next(i for i, h in enumerate(self.headers) if h.name.lower() == oldcolname.lower())
            new = change[4].split(" ")
            newcolname = new[-1]
            newcol = next(i for i, h in enumerate(self.headers) if h.name.lower() == newcolname.lower())
        except StopIteration:
            return False, "column_missing"
        if "n/a - Top ID" in change[3]:
            oldpar = ""
            oldpar_check = True
        else:
            oldpar = old[2]
            if oldpar.lower() not in self.nodes or self.nodes[oldpar.lower()].ps[oldcol] is None:
                oldpar_check = False
            else:
                oldpar_check = True
        if "n/a - Top ID" in change[4]:
            newpar = ""
            newpar_check = True
        else:
            newpar = new[2]
            if newpar.lower() not in self.nodes or self.nodes[newpar.lower()].ps[newcol] is None:
                newpar_check = False
            else:
                newpar_check = True
        if (
            self.headers[oldcol].type_ == "Parent"
            and self.headers[newcol].type_ == "Parent"
            and oldpar_check
            and newpar_check
        ):
            oldpc = int(self.pc)
            self.pc = newcol
            out = self.cut_paste_children(oldpar, newpar, oldcol, snapshot=False)
            self.pc = int(oldpc)
            if out["ok"]:
                self._log_prefixed(
                    log,
                    "Imported change | Cut and paste children",
                    change[2],
                    change[3],
                    change[4],
                )
                return True, None
            return False, "parent_mismatch"
        return False, "parent_mismatch"

    def apply_sort_later(self):
        d = self.sort_later_dct
        if not d:
            return
        if d.get("new_parent"):
            npk, h = d["new_parent"]
            if npk in self.nodes:
                self.nodes[npk].cn[h] = self.sort_node_cn(self.nodes[npk].cn[h], h)
        if d.get("new_parent_of_parent"):
            gpk, h = d["new_parent_of_parent"]
            if gpk in self.nodes:
                self.nodes[gpk].cn[h] = self.sort_node_cn(self.nodes[gpk].cn[h], h)
        old_hier = d.get("old_hier")
        if old_hier is not None:
            for iid in d.get("old_parents_of_parents") or ():
                if iid in self.nodes:
                    self.nodes[iid].cn[old_hier] = self.sort_node_cn(self.nodes[iid].cn[old_hier], old_hier)
        self.sort_later_dct = None

    def prune_changelog(self, up_to):
        if up_to < 0 or up_to >= len(self.changelog):
            return fail("usage", "Nothing to prune   ")
        removed = self.changelog[: up_to + 1]
        at_open = self.changelog_at_open
        from_at = removed[0].at if removed else ""
        to_at = removed[-1].at if removed else ""
        del self.changelog[: up_to + 1]
        if self.changelog_at_open:
            self.changelog_at_open = max(0, self.changelog_at_open - len(removed))
        self._log("Pruned changelog", f"From: {from_at} To: {to_at}", "", "")
        return ok({"removed": removed, "changelog_at_open": at_open, "n": len(removed)})

    def import_changes(self, rows, *, file_opened=""):
        changes = [list(r) for r in rows]
        if not changes:
            return fail("invalid_format", "File contains no data   ")
        row_len = max(map(len, changes), default=0)
        if row_len != 5:
            return fail("invalid_format", "Invalid changelog format   ")
        equalize_sublist_lens(seq=changes, len_=row_len)
        self._ensure_snapshot("full sheet", self.snapshot_sheet)
        log = ChangeBuilder(origin=ORIGIN_IMPORT)
        result_rows = []
        applied = 0
        unnecessary = 0
        failed = 0
        for i, change in enumerate(changes):
            try:
                ok_row, reason = self._import_one_change(change, log)
            except Exception:
                ok_row, reason = False, "parse"
            result_rows.append({"index": i, "ok": ok_row, "reason": reason})
            if ok_row and reason is None:
                applied += 1
            elif ok_row and reason == "unnecessary":
                unnecessary += 1
            else:
                failed += 1
        if applied:
            src = os.path.basename(file_opened) if file_opened else ""
            log.summary(
                f"Imported {applied} changes from: {src}",
                f"Unsuccessful: {failed} Unnecessary: {unnecessary}",
            )
            self.commit_change(log)
            self.pc = int(self.hiers[0])
            self.rns = {r[self.ic].lower(): i for i, r in enumerate(self.data)}
        else:
            self.vs.pop()
        return ok(
            {
                "applied": applied,
                "unnecessary": unnecessary,
                "failed": failed,
                "rows": result_rows,
            }
        )


def _normalize_compare_heads(heads, row_len):
    if len(heads) < row_len:
        heads += list(repeat("", row_len - len(heads)))
    warnings = []
    tally_of_heads = defaultdict(lambda: -1)
    for coln in range(len(heads)):
        cell = heads[coln]
        if not cell:
            cell = f"MISSING_{coln + 1}"
            warnings.append(f" - Missing header in column #{coln + 1}")
        hk = cell.lower()
        tally_of_heads[hk] += 1
        if tally_of_heads[hk] > 0:
            orig = cell
            x = 1
            while hk in tally_of_heads:
                cell = f"{orig}_DUPLICATED_{x}"
                hk = cell.lower()
                x += 1
            tally_of_heads[hk] += 1
            warnings.append(f" - Duplicate header in column #{coln + 1}")
        heads[coln] = cell
    return heads, warnings


def read_table(path):
    if not os.path.isfile(path):
        return fail("file_not_found", f"File not found: {path}")
    suffix = os.path.splitext(path)[1].lower()
    if suffix not in _DATA_SUFFIXES:
        return fail("invalid_format", "File must be .xlsx, .xlsm, .xls, .csv, .tsv or .json   ")
    if suffix in (".csv", ".tsv"):
        try:
            with open(path, "r") as fh:
                rows = csv_str_x_data(fh.read())
        except Exception as error_msg:
            return fail("invalid_format", f"{error_msg}")
        return ok({"rows": rows, "program_data": None})
    if suffix == ".json":
        try:
            j = get_json_from_file(path)
        except Exception as error_msg:
            return fail("invalid_format", f"{error_msg}")
        if "program_data" in j:
            try:
                program_data = b32_x_dict(j["program_data"])
            except Exception as error_msg:
                return fail("invalid_format", f"{error_msg}")
            headers = [h["name"] for h in program_data["headers"]]
            rows = [headers] + list(program_data["records"])
            return ok({"rows": rows, "program_data": program_data})
        json_format = get_json_format(j)
        if not json_format:
            return fail("invalid_format", "Could not find data of correct format   ")
        rows, _row_len = json_to_sheet(
            j,
            format_=json_format[0],
            key=json_format[1],
            get_format=False,
            return_rowlen=True,
        )
        return ok({"rows": rows, "program_data": None})
    try:
        in_mem = bytes_io_wb(path)
        wb = load_workbook(in_mem, read_only=True, data_only=True)
    except Exception as error_msg:
        return fail("invalid_format", f"{error_msg}")
    try:
        if len(wb.sheetnames) < 1:
            return fail("invalid_format", "File contains no data   ")
        if "program_data" in wb.sheetnames:
            ws = wb["program_data"]
            ws.reset_dimensions()
            try:
                program_data = b32_x_dict(ws_x_program_data_str(ws))
            except Exception as error_msg:
                return fail("invalid_format", f"{error_msg}")
            headers = [h["name"] for h in program_data["headers"]]
            rows = [headers] + list(program_data["records"])
            return ok({"rows": rows, "program_data": program_data})
        ws = wb[wb.sheetnames[0]]
        ws.reset_dimensions()
        rows = ws_x_data(ws)
        return ok({"rows": rows, "program_data": None})
    finally:
        wb.close()


def _compare_parent_value(nodes, ik, h):
    pk = nodes[ik].ps[h]
    if pk is None:
        return None
    if pk == "":
        return ""
    return nodes[pk].name


def compare_files(path_a, path_b, *, id_a=None, parents_a=None, id_b=None, parents_b=None):
    side_a = read_table(path_a)
    if not side_a["ok"]:
        return side_a
    side_b = read_table(path_b)
    if not side_b["ok"]:
        return side_b
    flags_a = id_a is not None or parents_a is not None
    flags_b = id_b is not None or parents_b is not None
    pd_a = side_a["result"]["program_data"]
    pd_b = side_b["result"]["program_data"]
    if pd_a is not None and flags_a:
        return fail("program_data_present", "File already has app data   ")
    if pd_b is not None and flags_b:
        return fail("program_data_present", "File already has app data   ")
    rows_a = [list(r) for r in side_a["result"]["rows"]]
    rows_b = [list(r) for r in side_b["result"]["rows"]]
    if not rows_a or not rows_b:
        return fail("invalid_format", "File contains no data   ")
    if pd_a is not None:
        ic_a = int(pd_a["ic"])
        parents_a_use = [int(h) for h in pd_a["hiers"]]
    else:
        if id_a is None or not parents_a:
            return fail("need_columns", "Need ID and parent columns   ")
        ic_a = int(id_a)
        parents_a_use = list(parents_a)
    if pd_b is not None:
        ic_b = int(pd_b["ic"])
        parents_b_use = [int(h) for h in pd_b["hiers"]]
    else:
        if id_b is None or not parents_b:
            return fail("need_columns", "Need ID and parent columns   ")
        ic_b = int(id_b)
        parents_b_use = list(parents_b)
    row_len_a = equalize_sublist_lens(rows_a)
    row_len_b = equalize_sublist_lens(rows_b)
    heads_a, warn_a = _normalize_compare_heads(rows_a[0].copy(), row_len_a)
    heads_b, warn_b = _normalize_compare_heads(rows_b[0].copy(), row_len_b)
    rows_a[0] = heads_a
    rows_b[0] = heads_b
    sheet_a, nodes_a, warn_a, rns_a = TreeBuilder().build(
        input_sheet=rows_a,
        output_sheet=[],
        row_len=row_len_a,
        ic=ic_a,
        hiers=parents_a_use,
        nodes={},
        warnings=warn_a,
        rns={},
        add_warnings=True,
        skip_1st=True,
        compare=True,
        fix_associate=True,
        strip=False,
    )
    sheet_b, nodes_b, warn_b, rns_b = TreeBuilder().build(
        input_sheet=rows_b,
        output_sheet=[],
        row_len=row_len_b,
        ic=ic_b,
        hiers=parents_b_use,
        nodes={},
        warnings=warn_b,
        rns={},
        add_warnings=True,
        skip_1st=True,
        compare=True,
        fix_associate=True,
        strip=False,
    )
    parset_a = set(parents_a_use)
    parset_b = set(parents_b_use)
    pcold = defaultdict(list)
    for i, h in enumerate(heads_a):
        if i in parset_a:
            pcold[h].append(i)
    for i, h in enumerate(heads_b):
        if i in parset_b:
            pcold[h].append(i)
    detcold = defaultdict(list)
    skip_a = {ic_a} | parset_a
    skip_b = {ic_b} | parset_b
    for i, h in enumerate(heads_a):
        if i not in skip_a:
            detcold[h].append(i)
    for i, h in enumerate(heads_b):
        if i not in skip_b:
            detcold[h].append(i)
    matching_hrs = sorted((k for k, v in pcold.items() if len(v) > 1), key=sort_key)
    matching_details = sorted((k for k, v in detcold.items() if len(v) > 1), key=sort_key)
    only_a = []
    only_b = []
    moved = []
    hdset_a_p = {h for i, h in enumerate(heads_a) if i in parset_a}
    hdset_b_p = {h for i, h in enumerate(heads_b) if i in parset_b}
    for h in hdset_a_p:
        if h not in hdset_b_p:
            only_a.append({"name": h, "role": "parent"})
    for h in hdset_b_p:
        if h not in hdset_a_p:
            only_b.append({"name": h, "role": "parent"})
    for name, idxs in pcold.items():
        if len(idxs) > 1 and idxs[0] != idxs[1]:
            moved.append({"name": name, "index_a": idxs[0], "index_b": idxs[1]})
    hdset_a_d = {h for i, h in enumerate(heads_a) if i not in skip_a}
    hdset_b_d = {h for i, h in enumerate(heads_b) if i not in skip_b}
    for h in hdset_a_d:
        if h not in hdset_b_d:
            only_a.append({"name": h, "role": "detail"})
    for h in hdset_b_d:
        if h not in hdset_a_d:
            only_b.append({"name": h, "role": "detail"})
    for name, idxs in detcold.items():
        if len(idxs) > 1 and idxs[0] != idxs[1]:
            moved.append({"name": name, "index_a": idxs[0], "index_b": idxs[1]})
    ids_only_a = [nodes_a[ik].name for ik in nodes_a if ik not in nodes_b]
    ids_only_b = [nodes_b[ik].name for ik in nodes_b if ik not in nodes_a]
    parent_diffs = []
    detail_diffs = []
    for ik in nodes_a:
        if ik not in nodes_b:
            continue
        ID = nodes_a[ik].name
        for nx in matching_hrs:
            h1 = pcold[nx][0]
            h2 = pcold[nx][1]
            a_val = _compare_parent_value(nodes_a, ik, h1)
            b_val = _compare_parent_value(nodes_b, ik, h2)
            if a_val != b_val:
                parent_diffs.append({"id": ID, "column": nx, "a": a_val, "b": b_val})
        for nx in matching_details:
            c1 = sheet_a[rns_a[ik]][detcold[nx][0]]
            c2 = sheet_b[rns_b[ik]][detcold[nx][1]]
            if c1.lower() != c2.lower():
                detail_diffs.append({"id": ID, "column": nx, "a": c1, "b": c2})
    identical = not (only_a or only_b or moved or ids_only_a or ids_only_b or parent_diffs or detail_diffs)
    return ok(
        {
            "identical": identical,
            "warnings_a": list(warn_a),
            "warnings_b": list(warn_b),
            "headers": {
                "id_column": {
                    "a": {"index": ic_a, "name": heads_a[ic_a]},
                    "b": {"index": ic_b, "name": heads_b[ic_b]},
                },
                "only_a": only_a,
                "only_b": only_b,
                "moved": moved,
            },
            "ids_only_a": ids_only_a,
            "ids_only_b": ids_only_b,
            "parent_diffs": parent_diffs,
            "detail_diffs": detail_diffs,
        }
    )
